"""Excel から値・文章・表を取り出す（openpyxl。計算結果の値を読む）.

セルの値は表示形式どおりの文字列で比べる（devtools_common.excel.render_value）. 例えば表示形式 `000` の数値 1 は
「001」、`#,##0` の 1500 は「1,500」、日付は `yyyy/mm/dd` なら「2026/04/05」. 文字列で入力した「001」と、
表示形式でゼロ埋めした数値 1 は同じ値とみなす（Excel の画面で同じに見えるため）.

表・文章は空行で打ち切らない. 連続した空行が end_blank_rows 行（既定 2）続くか、次の節の始まり
（設定にある別のラベル、空行の後の 1 セルだけの行、同じ見出し行の繰り返し）で終わりとする.
読んだセルは ReadContext.covered に記録する（設定に無いセルの変更を拾うため）.

走査するのは値のあるセルだけ（シートの「使用範囲」全体ではない）. 書式だけのセルが末尾の行・列まで広がった
シートでも、値のある範囲だけを読む. 値のあるセルが MAX_CELLS 個を超えるシートは、そこまでの行だけを読み、
警告 sheet_too_large を出す（sheet_alerts）.

数式セルは、Excel が保存した計算結果（キャッシュ）を読む. ツール（openpyxl など）で作った Excel は計算結果を
持たないため、そのままでは数式セルが空になる. その場合は数式の文字列（例: =SUM(B4:B5)）を値として読み、
警告を出す（Excel で開いて保存し直すと計算結果で比べられる）.
"""

from __future__ import annotations

import math
import weakref
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from devtools_common.excel import (
    Merge,
    MergeMap,
    header_depth,
    header_names,
    looks_numeric,
    render_cell,
    require_openpyxl,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..domain.model import text

MAX_HEADER_ROWS = 3
END_BLANK_ROWS = 2  # 表・文章の終わりとみなす連続した空行の数（既定）
MAX_CELLS = 1_000_000  # 1 シートで読む値のあるセルの上限（超えた分の行は読まずに警告）


@dataclass
class ReadContext:
    """表・文章の読み方の設定と、読んだ結果の記録（1 シート分）.

    stop_labels: 表・文章の終わりとみなすラベル（同じシートの設定にある別の項目のラベル・表題・見出し）
    end_blank_rows: 連続した空行がこの数だけ続いたら終わり
    covered: 読んだセル（行, 列）. 設定に無いセルの比較から除く
    alerts: 警告（rule_id, 文）. notes: 注記
    """

    stop_labels: set[str] = field(default_factory=set)
    end_blank_rows: int = END_BLANK_ROWS
    covered: set[tuple[int, int]] = field(default_factory=set)
    alerts: list[tuple[str, str]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def cover(self, ws: Worksheet, row: int, col: int) -> None:
        m = _merges(ws).get(row, col)
        if m:
            self.covered.update((r, c) for r in range(m.r0, m.r1 + 1) for c in range(m.c0, m.c1 + 1))
        else:
            self.covered.add((row, col))

    def cover_rows(self, ws: Worksheet, first: int, last: int) -> None:
        for r in range(first, last + 1):
            for c in range(1, _sheet(ws).max_col + 1):
                self.cover(ws, r, c)


def open_book(path: Path) -> Any:
    return load_book(path)[0]


def _formula_text(cell: Any) -> str | None:
    if getattr(cell, "data_type", None) != "f" or cell.value is None:
        return None
    v = cell.value
    if isinstance(v, str):
        return v
    text_attr = getattr(v, "text", None)  # 配列数式（ArrayFormula）
    return text_attr if isinstance(text_attr, str) else str(v)


def _load_with_formula_fallback(path: Path) -> tuple[Any, list[str]]:
    """devtools_common.excel.load_workbook_with_formula_fallback と同じ結果を、値のあるセルだけを見て作る.

    共通版は数式を読み取り専用モードの iter_rows で探すため、書式だけのセルが末尾まで広がったシートでは
    使用範囲の全セル（例: 5 万行 × 16,384 列）を回って非常に遅い. ここでは通常モードで読み、ファイルにあるセルだけを見る.
    """
    openpyxl = require_openpyxl()
    if not path.is_file():
        raise FileNotFoundError(f"Excel が見つかりません: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path)
    missing: list[str] = []
    try:
        for fws in formulas.worksheets:
            if fws.title not in wb.sheetnames or not hasattr(fws, "_cells"):
                continue
            ws = wb[fws.title]
            for (r, c), cell in sorted(fws._cells.items()):
                f = _formula_text(cell)
                if f is None:
                    continue
                target = ws._cells.get((r, c))
                if target is None or target.value is None:
                    target = ws.cell(r, c)
                    target.value = f
                    missing.append(f"{fws.title}!{target.coordinate}")
    finally:
        formulas.close()
    return wb, missing


def load_book(path: Path) -> tuple[Any, list[str]]:
    """計算結果の値で読んだブックと、警告（計算結果が保存されていない数式セルがあったとき）."""
    wb, missing = _load_with_formula_fallback(path)
    if not missing:
        return wb, []
    examples = "、".join(missing[:3]) + (" など" if len(missing) > 3 else "")
    return wb, [
        f"{path.name}: 計算結果が保存されていない数式セルが {len(missing)} 個あります（{examples}）。"
        "これらは計算結果ではなく数式の文字列（例: =SUM(B4:B5)）で比べました。ツールで作った Excel によくあります。"
        "Excel で開いて保存し直すと、計算結果で比べられます"
    ]


@dataclass
class _Sheet:
    """シートの値のあるセルの索引（1 シートにつき 1 回だけ作る）.

    cells: (行, 列) → セル（行・列の順）. rows: 行 → 値のある列（昇順）
    max_row / max_col: 値のある範囲の右下（値が無ければ 0）
    merges: 結合の索引（値のある範囲で切り詰める. 列や行全体の結合でもメモリを使わない）
    alerts: 読み方の警告（上限を超えて読まなかった行など）
    """

    cells: dict[tuple[int, int], Any]
    rows: dict[int, list[int]]
    max_row: int
    max_col: int
    merges: MergeMap
    alerts: list[tuple[str, str]] = field(default_factory=list)


_SHEETS: weakref.WeakKeyDictionary[Any, _Sheet] = weakref.WeakKeyDictionary()


def _stored_cells(ws: Worksheet) -> list[tuple[tuple[int, int], Any]]:
    """値のあるセル. openpyxl が持っているセルだけを見る（ws.cell / iter_rows は無いセルを作るので使わない）."""
    store = getattr(ws, "_cells", None)
    if isinstance(store, dict):
        return [(k, c) for k, c in store.items() if c.value is not None]
    return [((c.row, c.column), c) for row in ws.iter_rows() for c in row if c.value is not None]  # pragma: no cover


def _index(ws: Worksheet) -> _Sheet:
    found = sorted(_stored_cells(ws), key=lambda kc: kc[0])
    alerts: list[tuple[str, str]] = []
    if len(found) > MAX_CELLS:
        cut = found[MAX_CELLS][0][0]  # この行から下は読まない（行の途中では切らない）
        n = len(found)
        found = [kc for kc in found if kc[0][0] < cut]
        alerts.append(
            (
                "sheet_too_large",
                f"{ws.title}: 値のあるセルが {n:,} 個あり、上限 {MAX_CELLS:,} 個を超えたため、行 {cut - 1} までだけを読みました"
                f"（行 {cut} から下は比べていません）。シートを分けるか、不要な値を消してください",
            )
        )
    cells = dict(found)
    rows: dict[int, list[int]] = {}
    for r, c in cells:
        rows.setdefault(r, []).append(c)
    max_row = max(rows, default=0)
    max_col = max((cols[-1] for cols in rows.values()), default=0)
    return _Sheet(cells, rows, max_row, max_col, MergeMap(_clipped_merges(ws, max_row, max_col)), alerts)


def _clipped_merges(ws: Worksheet, max_row: int, max_col: int) -> Iterator[Merge]:
    """結合範囲を値のある範囲で切り詰める（左上が範囲外の結合は値を持たないので除く）."""
    ranges = getattr(getattr(ws, "merged_cells", None), "ranges", ()) or ()
    for r in ranges:
        if r.min_row > max_row or r.min_col > max_col:
            continue
        yield Merge(r.min_row, r.min_col, min(r.max_row, max(max_row, r.min_row)), min(r.max_col, max(max_col, r.min_col)))


def _sheet(ws: Worksheet) -> _Sheet:
    s = _SHEETS.get(ws)
    if s is None:
        s = _SHEETS[ws] = _index(ws)
    return s


def _merges(ws: Worksheet) -> MergeMap:
    return _sheet(ws).merges


def book_alerts(wb: Any, label: str) -> list[tuple[str, str]]:
    """読んだシートの警告（rule_id, 文）. 例: sheet_too_large（値のあるセルが上限を超えた）. label はファイル名など."""
    return [(rule, f"{label}: {msg}") for ws in wb.worksheets if ws in _SHEETS for rule, msg in _SHEETS[ws].alerts]


def cell_text(ws: Worksheet, row: int, col: int) -> str:
    """セルの値を表示形式どおりの文字列で（前後の空白は除く）. 結合で埋めない. 値の無いセルは作らずに空."""
    cell = _sheet(ws).cells.get((row, col))
    return "" if cell is None else text(render_cell(cell)[0])


def _merged_value(ws: Worksheet, row: int, col: int) -> str:
    """結合の左上の値で埋めたセルの値（表示形式どおりの文字列）."""
    return cell_text(ws, *_merges(ws).anchor(row, col))


def norm_label(v: Any) -> str:
    """ラベルの照合用に正規化する（末尾の「:」「：」、空白、セル内改行を無視）."""
    return _norm_label(v)


def outside_rows(ws: Worksheet, covered: set[tuple[int, int]]) -> list[tuple[int, dict[int, str]]]:
    """読んだセル（covered）以外の、値のあるセルを行ごとに（行番号, 列 → 値）. 結合の左上以外は除く."""
    rows: dict[int, dict[int, str]] = {}
    merges = _merges(ws)
    for (r, col), c in _sheet(ws).cells.items():
        if (r, col) in covered or merges.is_covered(r, col):
            continue
        v = text(render_cell(c)[0])
        if v:
            rows.setdefault(r, {})[col] = v
    return sorted(rows.items())


def _norm_label(v: Any) -> str:
    return text(v).rstrip(":：").replace("\n", "").replace(" ", "").replace("　", "")


def find_label(ws: Worksheet, label: str) -> tuple[int, int] | None:
    want = _norm_label(label)
    for pos, c in _sheet(ws).cells.items():  # 行・列の順
        if _norm_label(c.value) == want:
            return pos
    return None


def read_label(ws: Worksheet, label: str, offset: tuple[int, int] = (0, 1), ctx: ReadContext | None = None) -> str | None:
    """ラベルのセルから offset（行, 列）離れたセルの値. ラベルが無ければ None."""
    pos = find_label(ws, label)
    if pos is None:
        return None
    r, c = pos[0] + offset[0], pos[1] + offset[1]
    if ctx is not None:
        ctx.cover(ws, *pos)
        if r >= 1 and c >= 1:
            ctx.cover(ws, r, c)
    return _merged_value(ws, r, c)


def _row_cells(ws: Worksheet, r: int) -> list[tuple[int, str]]:
    """行 r の値のあるセル（列, 結合で埋めない値）."""
    out = []
    for c in _sheet(ws).rows.get(r, ()):
        if not _merges(ws).is_covered(r, c):
            v = cell_text(ws, r, c)
            if v:
                out.append((c, v))
    return out


def _starts_label(ws: Worksheet, r: int, ctx: ReadContext) -> bool:
    """行 r の左端の値が、設定にある別の項目のラベル（次の節の始まり）."""
    cells = _row_cells(ws, r)
    return bool(cells) and _norm_label(cells[0][1]) in ctx.stop_labels


def read_block(ws: Worksheet, label: str, offset: tuple[int, int] = (1, 0), ctx: ReadContext | None = None) -> str | None:
    """ラベルの下に続く行を改行でつないだ文章.

    途中の空行は空の行として含める. 連続した空行が end_blank_rows 行続くか、
    設定にある別のラベルの行、空行の後で他の列にも値がある行（表の始まり）で終わる.
    """
    ctx = ctx if ctx is not None else ReadContext()
    pos = find_label(ws, label)
    if pos is None:
        return None
    ctx.cover(ws, *pos)
    r, c = pos[0] + offset[0], pos[1] + offset[1]
    lines: list[str] = []
    blank = 0
    gaps: list[int] = []
    while r <= _sheet(ws).max_row:
        v = _merged_value(ws, r, c)
        if not v:
            blank += 1
            if blank >= ctx.end_blank_rows:
                break
            r += 1
            continue
        if _starts_label(ws, r, ctx) and lines:
            break
        if blank and (not lines or any(col != c for col, _ in _row_cells(ws, r))):
            break  # 空行の後に別の列の値がある行 = 次の表・節
        if blank:
            gaps += list(range(r - blank, r))
            lines += [""] * blank
        blank = 0
        lines.append(v)
        ctx.cover(ws, r, c)
        r += 1
    if gaps:
        ctx.notes.append(f"{ws.title}: 文章「{label}」は空行（行 {_rows(gaps)}）を挟んで続く行も含めて比べました")
    return "\n".join(lines)


def _rows(rows: list[int]) -> str:
    return "、".join(str(r) for r in rows[:5]) + (" など" if len(rows) > 5 else "")


def _own_value(ws: Worksheet, row: int, col: int) -> str:
    """結合の左上以外は空として読む（結合で埋めない）セルの値."""
    return "" if _merges(ws).is_covered(row, col) else _norm_label(cell_text(ws, row, col))


def _count_header_rows(ws: Worksheet, head_row: int) -> int:
    """見出しが何段か（結合セルから判断する. devtools_common.excel.header_depth）.

    次の行も見出しとみなすのは、次のどちらかのとき:
    - 見出し行から始まる結合セルが次の行まで縦に続いていて、次の行が文字だけ（例: 縦に結合した「項目名」の横に「型」「桁」）
    - 見出し行で横に結合したセル（例: 「型情報」）の下の各列に、どれも文字の値がある（例: 「型」「桁」）
    データ行も同じように横に結合しているだけの表や、結合の下が数値の表は 1 段のままです.
    """
    return header_depth(
        _merges(ws),
        lambda r, c: _own_value(ws, r, c),
        head_row,
        first_col=1,
        last_col=_sheet(ws).max_col,
        last_row=_sheet(ws).max_row,
        max_rows=MAX_HEADER_ROWS,
    )


def read_table(
    ws: Worksheet,
    header: str,
    title: str | None = None,
    header_rows: int | None = None,
    ctx: ReadContext | None = None,
) -> tuple[list[str], list[dict[str, str]]] | None:
    """見出しの列名 header を含む行を見出し行とする表. title があれば、そのラベルの次の行を見出し行とする.

    見出しが複数段（結合したグループ見出し + 子見出し）のときは「親/子」の列名にする.
    段数は header_rows で指定する. 省略時は結合セルから判断する（_count_header_rows）.
    同じ列名が 2 列以上あれば、2 列目から「備考#2」のように番号を付けて両方読む（警告を出す）.
    行は途中の空行を飛ばして読み、連続した空行が end_blank_rows 行続くか、次の節の始まりで終わる.
    結合セルは左上の値で埋める.
    """
    ctx = ctx if ctx is not None else ReadContext()
    if title:
        pos = find_label(ws, title)
        if pos is None:
            return None
        head_row = pos[0] + 1
        ctx.cover(ws, *pos)
    else:
        pos = find_label(ws, header)
        if pos is None:
            return None
        head_row = pos[0]
    depth = header_rows or _count_header_rows(ws, head_row)
    width = range(1, _sheet(ws).max_col + 1)
    names = header_names(lambda r, c: _norm_label(_merged_value(ws, r, c)), head_row, depth, width)
    cols: list[tuple[int, str]] = []
    seen: dict[str, list[int]] = {}
    prev = ""
    for c, name in zip(width, names, strict=True):
        same_merge = name == prev and all(_merges(ws).anchor(r, c)[1] < c for r in range(head_row, head_row + depth))
        prev = name
        if not name or same_merge:  # 横に結合した見出しの 2 列目以降は同じ列（データも結合で埋まる）
            continue
        seen.setdefault(name, []).append(c)
        n = len(seen[name])
        cols.append((c, name if n == 1 else f"{name}#{n}"))
    if not cols:
        return None
    where = title or header
    for name, cs in seen.items():
        if len(cs) > 1:
            letters = "・".join(f"{get_column_letter(c)} 列" for c in cs)
            renamed = "・".join(f"「{name}#{i}」" for i in range(2, len(cs) + 1))
            ctx.alerts.append(
                (
                    "duplicate_column",
                    f"{ws.title}: 表「{where}」に同じ列名「{name}」が {len(cs)} 列あります（{letters}）。"
                    f"2 列目から {renamed} として比べました",
                )
            )
    header_set = {n for _, n in cols}
    rows: list[dict[str, str]] = []
    r = head_row + depth
    last = head_row + depth - 1
    blank = 0
    gaps: list[int] = []
    while r <= _sheet(ws).max_row:
        values = {name: _merged_value(ws, r, c) for c, name in cols}
        if not any(values.values()):
            blank += 1
            if blank >= ctx.end_blank_rows:
                _warn_if_truncated(ws, r + 1, cols, rows, ctx, where)
                break
            r += 1
            continue
        reason = _section_start(ws, r, ctx, header_set, after_blank=blank > 0)
        if reason == "single_cell":
            _warn_single_cell_end(ws, r, ctx, where)
        if reason:
            break
        if blank:
            gaps += list(range(r - blank, r))
        blank = 0
        rows.append(values)
        last = r
        r += 1
    ctx.cover_rows(ws, head_row, last)
    if gaps:
        ctx.notes.append(f"{ws.title}: 表「{where}」は空行（行 {_rows(gaps)}）を挟んで続く行も表として比べました")
    return [n for _, n in cols], rows


def _section_start(ws: Worksheet, r: int, ctx: ReadContext, header_set: set[str], *, after_blank: bool) -> str:
    """行 r が次の節の始まり（表の終わり）なら理由、そうでなければ "".

    "label": 設定にある別の項目のラベル / "header": 同じ見出し行の繰り返し（次の表）/
    "single_cell": 空行の後の 1 セルだけの行（次の節の表題とみなす. 設定に無いので推測）
    """
    if _starts_label(ws, r, ctx):
        return "label"
    cells = _row_cells(ws, r)
    if {_norm_label(v) for _, v in cells} <= header_set and len(cells) >= 2:
        return "header"
    if after_blank and len(cells) == 1 and len(header_set) >= 2:
        return "single_cell"
    return ""


def _warn_single_cell_end(ws: Worksheet, r: int, ctx: ReadContext, where: str) -> None:
    """空行の後の 1 セルだけの行で表を終わりにしたとき: 表の続き（1 列だけ書いた行）かもしれないので警告する."""
    (c, v), *_ = _row_cells(ws, r)
    ctx.alerts.append(
        (
            "table_truncated",
            f"{ws.title}: 表「{where}」は空行の後の 1 セルだけの行 {r}（{get_column_letter(c)}{r}「{v}」）を次の節の表題とみなして"
            "終わりとしました（その行から下は表として比べていません）。次の節の表題なら、その表も設定に書くとこの警告は出ません。"
            "表の続きなら空行を詰めてください",
        )
    )


def _warn_if_truncated(
    ws: Worksheet, start: int, cols: list[tuple[int, str]], rows: list[dict[str, str]], ctx: ReadContext, where: str
) -> None:
    """表の終わりと判断した空行の下に、表の続きらしい行があれば警告する."""
    for r in range(start, _sheet(ws).max_row + 1):
        values = [_merged_value(ws, r, c) for c, _ in cols]
        if not any(values):
            continue
        if _section_start(ws, r, ctx, {n for _, n in cols}, after_blank=True):
            return
        filled = sum(1 for v in values if v)
        if filled < max(2, math.ceil(len(cols) * 0.5)):
            return
        first = [row[cols[0][1]] for row in rows if row.get(cols[0][1])]
        if first and all(looks_numeric(v) for v in first) and not looks_numeric(values[0]):
            return  # 1 列目の型が違う（承認者欄など）
        ctx.alerts.append(
            (
                "table_truncated",
                f"{ws.title}: 表「{where}」は空行が {ctx.end_blank_rows} 行続いたので終わりとしましたが、"
                f"下の行 {r} にも表の続きらしい値があります（比べていません）。"
                "表の続きなら設定の end_blank_rows を増やしてください",
            )
        )
        return
