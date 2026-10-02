"""サンプルの設計書（架空の会員・注文システム）を生成する.

    python samples/make_samples.py            # samples/ の下に書き出す
    python samples/make_samples.py -o DIR     # 別の場所へ

生成するもの:
- before/ と after/ : 新旧の API 設計書（API0001）とバッチ設計書（JOB0001）。差分の種類が一通り出るように変えてある
- spec/API定義書.xlsx と spec/openapi.yaml : 整合チェック用。わざと不一致を入れてある
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEAD = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)

API_V1: dict[str, Any] = {
    "id": "API0001",
    "name": "会員照会",
    "method": "GET",
    "path": "/api/v1/members/{memberId}",
    "summary": ["会員 ID で会員を 1 件取得する。", "存在しない場合は APPE0001 を返す。"],
    "request": [
        ["memberId", "文字列", "○", 10, "path", "会員 ID"],
    ],
    "response": [
        ["memberId", "文字列", "○", 10, "会員 ID"],
        ["name", "文字列", "○", 50, "氏名"],
        ["email", "文字列", "", 100, "メールアドレス"],
        ["status", "文字列", "○", 1, "状態（1: 有効 / 9: 退会）"],
        ["faxNumber", "文字列", "", 15, "FAX 番号"],
    ],
    "statuses": [["200", "正常"], ["404", "会員が存在しない（APPE0001）"]],
}

API_V2: dict[str, Any] = copy.deepcopy(API_V1)
API_V2["name"] = "会員詳細照会"  # 値の変更
API_V2["summary"].append("退会済みの会員も返す（status = 9）。")  # 文章の変更
API_V2["request"].append(["includeHistory", "真偽値", "", "", "query", "履歴を含めるか"])  # 行の追加
API_V2["response"][1][3] = 60  # セルの変更（桁）
API_V2["response"] = [r for r in API_V2["response"] if r[0] != "faxNumber"]  # 行の削除
API_V2["response"].append(["createdAt", "日時", "○", "", "登録日時"])  # 行の追加
API_V2["statuses"].append(["400", "入力誤り（APPE0002）"])

API0002: dict[str, Any] = {
    "id": "API0002",
    "name": "会員登録",
    "method": "POST",
    "path": "/api/v1/members",
    "summary": ["会員を 1 件登録する。"],
    "request": [
        ["name", "文字列", "○", 50, "body", "氏名"],
        ["email", "文字列", "○", 100, "body", "メールアドレス"],
    ],
    "response": [
        ["memberId", "文字列", "○", 10, "採番した会員 ID"],
    ],
    "statuses": [["201", "登録した"], ["400", "入力誤り（APPE0002）"]],
}

HISTORY_V1 = [["1.0", "2026-04-01", "初版"]]
HISTORY_V2 = [*HISTORY_V1, ["1.1", "2026-09-15", "退会会員の扱いを追加"]]

JOB_V1: dict[str, Any] = {
    "id": "JOB0001",
    "name": "注文ファイル取込",
    "trigger": "毎日 02:00（スケジューラ）",
    "summary": ["ストレージの注文 CSV を取り込み、注文テーブルに登録する。", "既にある注文番号は登録しない（冪等）。"],
    "steps": [
        ["1", "注文 CSV を取得する", "ストレージ", "作業ファイル"],
        ["2", "1 行ずつ検証する", "作業ファイル", "-"],
        ["3", "注文を登録する", "-", "注文テーブル"],
    ],
    "codes": [["0", "正常終了"], ["8", "異常終了"]],
}
JOB_V2: dict[str, Any] = copy.deepcopy(JOB_V1)
JOB_V2["trigger"] = "毎日 01:30（スケジューラ）"
JOB_V2["steps"][1][1] = "1 行ずつ検証し、誤りの行は読み飛ばす"
JOB_V2["steps"].append(["4", "結果ファイルを出力する", "-", "結果ファイル"])
JOB_V2["codes"].insert(1, ["4", "警告終了（読み飛ばした行あり）"])


def _title(ws, text: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)


def _label(ws, row: int, label: str, value: Any) -> None:
    ws.cell(row, 1, label).font = BOLD
    ws.cell(row, 1).fill = HEAD
    ws.cell(row, 2, value)


def _table(ws, row: int, title: str | None, header: list[str], rows: list[list[Any]]) -> int:
    """表を書き、次に使える行（空行を 1 行あけた先）を返す."""
    if title:
        ws.cell(row, 1, title).font = BOLD
        row += 1
    for c, h in enumerate(header, start=1):
        ws.cell(row, c, h).font = BOLD
        ws.cell(row, c).fill = HEAD
    for r in rows:
        row += 1
        for c, v in enumerate(r, start=1):
            ws.cell(row, c, None if v == "" else v)
    return row + 2


def api_sheet(ws, api: dict[str, Any]) -> None:
    _title(ws, "API 設計書")
    for i, (label, key) in enumerate([("API ID", "id"), ("API 名", "name"), ("メソッド", "method"), ("パス", "path")]):
        _label(ws, 3 + i, label, api[key])
    ws.cell(8, 1, "処理概要").font = BOLD
    for i, line in enumerate(api["summary"]):
        ws.cell(9 + i, 1, line)
    row = 9 + len(api["summary"]) + 1
    row = _table(ws, row, "入力項目", ["項目名", "型", "必須", "桁", "位置", "説明"], api["request"])
    row = _table(ws, row, "出力項目", ["項目名", "型", "必須", "桁", "説明"], api["response"])
    _table(ws, row, "応答", ["ステータス", "説明"], api["statuses"])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30


def api_book(apis: list[dict[str, Any]], history: list[list[str]] | None) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    if history:
        ws = wb.create_sheet("改訂履歴")
        _title(ws, "改訂履歴")
        _table(ws, 3, None, ["版", "日付", "内容"], history)
    for api in apis:
        api_sheet(wb.create_sheet(f"{api['id']}_{api['name']}"), api)
    return wb


def job_book(job: dict[str, Any]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    _title(ws, "バッチ設計書")
    _label(ws, 3, "ジョブ ID", job["id"])
    _label(ws, 4, "ジョブ名", job["name"])
    _label(ws, 5, "起動条件", job["trigger"])
    ws.cell(7, 1, "処理概要").font = BOLD
    for i, line in enumerate(job["summary"]):
        ws.cell(8 + i, 1, line)
    _table(wb.create_sheet("処理詳細"), 1, None, ["STEP", "処理", "入力", "出力"], job["steps"])
    _table(wb.create_sheet("終了コード"), 1, None, ["コード", "意味"], job["codes"])
    return wb


OPENAPI = """\
openapi: 3.0.3
info:
  title: 会員 API（サンプル）
  version: 1.0.0
paths:
  /api/v1/members/{memberId}:
    get:
      x-api-id: API0001
      operationId: getMember
      parameters:
        - name: memberId
          in: path
          required: true
          schema: {type: string, maxLength: 10}
        - name: includeHistory
          in: query
          schema: {type: boolean}
      responses:
        "200":
          description: 正常
          content:
            application/json:
              schema: {$ref: "#/components/schemas/Member"}
        "404": {description: 会員が存在しない}
  /api/v1/members:
    post:
      x-api-id: API0002
      operationId: registerMember
      requestBody:
        content:
          application/json:
            schema:
              type: object
              required: [name]
              properties:
                name: {type: string, maxLength: 50}
                email: {type: string, maxLength: 100}
                phone: {type: string, maxLength: 15}
      responses:
        "201":
          description: 登録した
          content:
            application/json:
              schema:
                type: object
                required: [memberId]
                properties:
                  memberId: {type: string, maxLength: 10}
        "400": {description: 入力誤り}
  /api/v1/members/{memberId}/withdraw:
    post:
      x-api-id: API0003
      operationId: withdrawMember
      parameters:
        - {name: memberId, in: path, required: true, schema: {type: string, maxLength: 10}}
      responses:
        "204": {description: 退会した}
components:
  schemas:
    Member:
      type: object
      required: [memberId, name, status, createdAt]
      properties:
        memberId: {type: string, maxLength: 10}
        name: {type: string, maxLength: 50}
        email: {type: string, maxLength: 100}
        status: {type: string, maxLength: 1}
        createdAt: {type: string, format: date-time}
"""


def generate(out: Path) -> list[Path]:
    written: list[Path] = []
    for sub, api, history, job in (("before", API_V1, HISTORY_V1, JOB_V1), ("after", API_V2, HISTORY_V2, JOB_V2)):
        d = out / sub
        d.mkdir(parents=True, exist_ok=True)
        for name, wb in ((f"{api['id']}_会員照会.xlsx", api_book([api], history)), ("JOB0001_注文取込.xlsx", job_book(job))):
            wb.save(d / name)
            written.append(d / name)
    d = out / "spec"
    d.mkdir(parents=True, exist_ok=True)
    api_book([API_V2, API0002], None).save(d / "API定義書.xlsx")
    (d / "openapi.yaml").write_text(OPENAPI, encoding="utf-8", newline="\n")
    written += [d / "API定義書.xlsx", d / "openapi.yaml"]
    return written


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--out", type=Path, default=Path(__file__).resolve().parent)
    for p in generate(ap.parse_args().out):
        print(p)


if __name__ == "__main__":
    main()
