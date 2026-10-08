"""採点で見つかった取りこぼしの回帰テスト.

- 同じ名前の列（「備考」が 2 列）の 2 列目も比べる
- 表・文章の途中の空行で打ち切らない（空行より下の変更も出す）。打ち切ったら警告
- 設定にないセル（承認者欄など）の変更も「比較対象外のセルの変更」として出す
- キーが重複した行も出現順で対応付けて比べる
- --all: -f・-o に従う、出力名の衝突は番号を付ける、欠けたペアは警告して続ける、--config-dir はサブコマンドの後でもよい
"""

from __future__ import annotations

import contextlib
import io
import json
import time
from pathlib import Path

import pytest
import yaml
from devtools_common.cli import run_main
from openpyxl import Workbook

from design_diff.application import excel_diff, render
from design_diff.cli import main as cli
from design_diff.domain.excel_diff import compare_table


def book(path: Path, cells: dict[str, object], title: str = "IF0001") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for ref, v in cells.items():
        ws[ref] = v
    wb.save(path)
    return path


def table_cells(header: list[str], rows: list[list[object]], start: int = 1) -> dict[str, object]:
    cells: dict[str, object] = {}
    for r, row in enumerate([header, *rows], start=start):
        for c, v in enumerate(row):
            if v is not None:
                cells[f"{chr(ord('A') + c)}{r}"] = v
    return cells


TABLE_CONF = {"sheets": [{"name": "IF*", "items": [{"kind": "table", "name": "レイアウト", "header": "項番", "key": ["項番"]}]}]}


def changes_of(conf, before: Path, after: Path, **kw):
    return [(c.category, c.field, c.before, c.after) for c in excel_diff.run(conf, before, after, **kw).changes]


# --- 同じ名前の列 -----------------------------------------------------------------------------


def test_duplicate_column_names_are_numbered_and_compared(tmp_path: Path):
    header = ["項番", "備考", "型", "備考"]
    before = book(tmp_path / "b.xlsx", table_cells(header, [[1, "a", "数値", "初版"], [2, "b", "文字列", "x"]]))
    after = book(tmp_path / "a.xlsx", table_cells(header, [[1, "a", "数値", "改訂"], [2, "b", "文字列", "x"]]))
    result = excel_diff.run(TABLE_CONF, before, after)
    assert [(c.category, c.field, c.before, c.after) for c in result.changes] == [("cell_changed", "1 の 備考#2", "初版", "改訂")]
    assert any(rule == "duplicate_column" and "備考#2" in msg for rule, msg in result.alerts)


# --- 空行をはさむ表・文章 -----------------------------------------------------------------------


def test_rows_after_a_blank_row_are_compared(tmp_path: Path):
    header = ["項番", "項目名", "桁"]
    rows_b = [[1, "id", 10], [2, "name", 50], [None, None, None], [3, "email", 100]]
    rows_a = [[1, "id", 10], [2, "name", 50], [None, None, None], [3, "email", 200]]
    before = book(tmp_path / "b.xlsx", table_cells(header, rows_b))
    after = book(tmp_path / "a.xlsx", table_cells(header, rows_a))
    result = excel_diff.run(TABLE_CONF, before, after)
    assert [(c.category, c.field, c.before, c.after) for c in result.changes] == [("cell_changed", "3 の 桁", "100", "200")]
    assert any("空行" in n for n in result.notes)  # 空行を挟んで読んだことを知らせる


def test_table_stops_at_next_section_title(tmp_path: Path):
    cells = table_cells(["項番", "項目名"], [[1, "id"], [2, "name"]], start=2)
    cells |= {"A1": "入力項目", "A5": "出力項目"} | table_cells(["項番", "項目名"], [[1, "result"]], start=6)
    conf = {
        "sheets": [
            {
                "name": "IF*",
                "items": [
                    {"kind": "table", "name": "入力", "title": "入力項目", "header": "項番", "key": ["項番"]},
                    {"kind": "table", "name": "出力", "title": "出力項目", "header": "項番", "key": ["項番"]},
                ],
            }
        ]
    }
    before = book(tmp_path / "b.xlsx", cells)
    after = book(tmp_path / "a.xlsx", cells | {"B7": "code"})
    assert changes_of(conf, before, after) == [("cell_changed", "1 の 項目名", "result", "code")]


def test_truncated_table_is_warned(tmp_path: Path):
    header = ["項番", "項目名", "桁"]
    rows = [[1, "id", 10], [None] * 3, [None] * 3, [None] * 3, [2, "name", 50]]
    before = book(tmp_path / "b.xlsx", table_cells(header, rows))
    result = excel_diff.run(TABLE_CONF, before, before)
    assert any(rule == "table_truncated" and "行 6" in msg for rule, msg in result.alerts)


def test_block_continues_after_a_blank_row(tmp_path: Path):
    conf = {"sheets": [{"name": "IF*", "items": [{"kind": "block", "name": "処理概要", "label": "処理概要"}]}]}
    base = {"A1": "処理概要", "A2": "1 行目", "A4": "3 行目（空行の後）"}
    before = book(tmp_path / "b.xlsx", base)
    after = book(tmp_path / "a.xlsx", base | {"A4": "3 行目を変えた"})
    (cat, _, b, a), *rest = changes_of(conf, before, after)
    assert (cat, rest) == ("value_changed", [])
    assert b.endswith("3 行目（空行の後）") and a.endswith("3 行目を変えた")


def test_end_blank_rows_option(tmp_path: Path):
    header = ["項番", "項目名", "桁"]
    rows_b = [[1, "id", 10], [None] * 3, [None] * 3, [2, "name", 50]]
    rows_a = [[1, "id", 10], [None] * 3, [None] * 3, [2, "name", 60]]
    before = book(tmp_path / "b.xlsx", table_cells(header, rows_b))
    after = book(tmp_path / "a.xlsx", table_cells(header, rows_a))
    conf = {"sheets": [{"name": "IF*", "items": [dict(TABLE_CONF["sheets"][0]["items"][0], end_blank_rows=3)]}]}
    assert changes_of(conf, before, after) == [("cell_changed", "2 の 桁", "50", "60")]


# --- 設定にないセル -----------------------------------------------------------------------------

LABEL_CONF = {"sheets": [{"name": "IF*", "items": [{"kind": "label", "name": "IF ID", "label": "IF ID"}]}]}


def test_cells_outside_config_are_reported(tmp_path: Path):
    base = {"A1": "IF ID", "B1": "IF0001", "D1": "承認者", "E1": "山田"}
    before = book(tmp_path / "b.xlsx", base)
    after = book(tmp_path / "a.xlsx", base | {"E1": "佐藤"})
    (change,) = excel_diff.run(LABEL_CONF, before, after).changes
    assert change.category == "outside_changed"
    assert "1 件" in change.after and "E1" in change.after


def test_cells_outside_config_all_option(tmp_path: Path):
    base = {"A1": "IF ID", "B1": "IF0001", "D1": "承認者", "E1": "山田"}
    before = book(tmp_path / "b.xlsx", base)
    after = book(tmp_path / "a.xlsx", base | {"E1": "佐藤", "E2": "追記"})
    got = changes_of(LABEL_CONF, before, after, outside="all")
    assert got == [("outside_cell_changed", "E1", "山田", "佐藤"), ("outside_cell_changed", "E2", "", "追記")]
    assert changes_of(LABEL_CONF, before, after, outside="off") == []


def test_outside_cells_follow_inserted_rows(tmp_path: Path):
    """表の行が増えて下の承認者欄がずれても、内容が同じなら差分にしない."""
    rows_b = [[1, "id"], [2, "name"]]
    rows_a = [[1, "id"], [2, "name"], [3, "email"]]
    before = book(tmp_path / "b.xlsx", table_cells(["項番", "項目名"], rows_b) | {"A6": "承認者", "B6": "山田"})
    after = book(tmp_path / "a.xlsx", table_cells(["項番", "項目名"], rows_a) | {"A7": "承認者", "B7": "山田"})
    assert [c[0] for c in changes_of(TABLE_CONF, before, after)] == ["row_added"]


def test_sheet_without_config_is_compared(tmp_path: Path):
    def make(path: Path, approver: str) -> Path:
        wb = Workbook()
        wb.active.title = "IF0001"
        wb.active["A1"] = "IF ID"
        wb.active["B1"] = "IF0001"
        ws = wb.create_sheet("表紙")
        ws["B5"] = approver
        wb.save(path)
        return path

    changes = excel_diff.run(LABEL_CONF, make(tmp_path / "b.xlsx", "山田"), make(tmp_path / "a.xlsx", "佐藤")).changes
    assert [(c.category, c.location) for c in changes] == [("outside_changed", "表紙")]


# --- 重複したキー -------------------------------------------------------------------------------


def test_duplicate_keys_are_compared_in_order():
    cols = ["項番", "内容"]
    before = (cols, [{"項番": "1", "内容": "a"}, {"項番": "1", "内容": "b"}])
    after = (cols, [{"項番": "1", "内容": "a"}, {"項番": "1", "内容": "c"}])
    got = [(c.category, c.field, c.before, c.after) for c in compare_table("表", ["項番"], before, after)]
    assert ("cell_changed", "1（2 件目） の 内容", "b", "c") in got
    assert ("duplicate_key", "1", "", "") in got
    assert len(got) == 2


def test_duplicate_key_extra_row_is_added():
    cols = ["項番", "内容"]
    before = (cols, [{"項番": "1", "内容": "a"}])
    after = (cols, [{"項番": "1", "内容": "a"}, {"項番": "1", "内容": "b"}])
    got = [(c.category, c.field) for c in compare_table("表", ["項番"], before, after)]
    assert ("row_added", "1（2 件目）") in got


# --- --all --------------------------------------------------------------------------------------


def _all_setup(tmp_path: Path, pairs: list[dict]) -> Path:
    conf_dir = tmp_path / "conf"
    conf_dir.mkdir()
    (conf_dir / "if.yaml").write_text(yaml.safe_dump(TABLE_CONF, allow_unicode=True), encoding="utf-8")
    g = {"output_dir": str(tmp_path / "out"), "pairs": pairs}
    (conf_dir / "global.yaml").write_text(yaml.safe_dump(g, allow_unicode=True), encoding="utf-8")
    return conf_dir


def _run(*args: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue(), err.getvalue()


@pytest.fixture()
def two_books(tmp_path: Path) -> tuple[Path, Path, Path]:
    header = ["項番", "項目名"]
    b = book(tmp_path / "b.xlsx", table_cells(header, [[1, "id"]]))
    (tmp_path / "x").mkdir()
    (tmp_path / "y").mkdir()
    a1 = book(tmp_path / "x" / "spec.xlsx", table_cells(header, [[1, "code"]]))
    a2 = book(tmp_path / "y" / "spec.xlsx", table_cells(header, [[1, "key"]]))
    return b, a1, a2


def test_all_name_collision_gets_number(tmp_path: Path, two_books):
    b, a1, a2 = two_books
    conf_dir = _all_setup(
        tmp_path, [{"type": "if", "before": str(b), "after": str(a1)}, {"type": "if", "before": str(b), "after": str(a2)}]
    )
    rc, _, _ = _run("excel", "--all", "--config-dir", str(conf_dir))  # --config-dir はサブコマンドの後でもよい
    assert rc == 0
    names = sorted(p.name for p in (tmp_path / "out").iterdir())
    assert names == ["diff_spec-2.md", "diff_spec.md"]
    assert "code" in (tmp_path / "out" / "diff_spec.md").read_text(encoding="utf-8")
    assert "key" in (tmp_path / "out" / "diff_spec-2.md").read_text(encoding="utf-8")


def test_all_format_and_output(tmp_path: Path, two_books):
    b, a1, _ = two_books
    conf_dir = _all_setup(tmp_path, [{"type": "if", "before": str(b), "after": str(a1)}])
    summary = tmp_path / "summary.json"
    rc, _, _ = _run("--config-dir", str(conf_dir), "excel", "--all", "-f", "json", "-o", str(summary))
    assert rc == 0
    data = json.loads((tmp_path / "out" / "diff_spec.json").read_text(encoding="utf-8"))
    assert data["findings"][0]["rule_id"] == "cell_changed"
    s = json.loads(summary.read_text(encoding="utf-8"))
    assert s["meta"]["pairs"] == 1 and s["findings"][0]["rule_id"] == "pair_compared"


def test_all_missing_pair_continues(tmp_path: Path, two_books):
    b, a1, _ = two_books
    pairs = [{"type": "if", "before": str(b), "after": str(tmp_path / "none.xlsx")}, {"type": "if", "before": str(b), "after": str(a1)}]
    conf_dir = _all_setup(tmp_path, pairs)
    rc, out, err = _run("excel", "--all", "--config-dir", str(conf_dir))
    assert rc == 1  # 欠けたペアがあれば最後に 1
    assert (tmp_path / "out" / "diff_spec.md").is_file()  # 残りのペアは比べる
    assert "none.xlsx" in err and "none.xlsx" in out


def test_config_dir_after_subcommand_for_spec(samples_dir: Path):
    rc, out, _ = _run(
        "spec",
        "--excel",
        str(samples_dir / "spec/API定義書.xlsx"),
        "--oas",
        str(samples_dir / "spec/openapi.yaml"),
        "--config-dir",
        str(Path(__file__).resolve().parent.parent / "config"),
    )
    assert rc in (0, 1) and "API" in out


@pytest.fixture()
def samples_dir(tmp_path: Path) -> Path:
    import importlib.util

    root = Path(__file__).resolve().parent.parent
    spec = importlib.util.spec_from_file_location("make_samples_r", root / "samples/make_samples.py")
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.generate(tmp_path / "s")
    return tmp_path / "s"


# --- 第 2 回: 読む範囲・pair_by の重複・1 セル行での打ち切り・global.yaml の相対パス・python -m ----------------


def test_format_only_cells_far_away_are_not_scanned(tmp_path: Path):
    """書式だけのセルが末尾まで広がっても、値のある範囲だけを読む（使用範囲全体のセルを作らない）."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    from design_diff.infrastructure import excel_reader

    p = book(tmp_path / "b.xlsx", table_cells(["項番", "項目名"], [[1, "id"], [2, "name"]]))
    wb = load_workbook(p)
    wb.active["XFD50000"].font = Font(bold=True)  # 書式だけ（値なし）
    wb.active.merge_cells("D1:XFD1")  # 値の無い列まで続く結合
    wb.save(p)
    started = time.perf_counter()
    wb = excel_reader.open_book(p)
    ws = wb.active
    assert ws.max_row == 50000  # openpyxl の使用範囲は書式のセルまで
    stored = len(ws._cells)  # 結合の覆われたセルは openpyxl が読み込み時に作る
    cols, rows = excel_reader.read_table(ws, "項番")
    assert cols == ["項番", "項目名"] and len(rows) == 2
    excel_reader.outside_rows(ws, set())
    assert len(ws._cells) == stored  # 走査で空のセルを作っていない
    assert changes_of(TABLE_CONF, p, p) == []
    assert time.perf_counter() - started < 30  # 使用範囲（5 万行 × 16,384 列）を回ると数分かかる


def test_sheet_over_cell_limit_is_warned(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from design_diff.infrastructure import excel_reader

    monkeypatch.setattr(excel_reader, "MAX_CELLS", 5)
    header = ["項番", "項目名"]
    before = book(tmp_path / "b.xlsx", table_cells(header, [[1, "id"], [2, "name"], [3, "x"]]))
    after = book(tmp_path / "a.xlsx", table_cells(header, [[1, "id"], [2, "name"], [3, "y"]]))
    result = excel_diff.run(TABLE_CONF, before, after)
    alerts = [msg for rule, msg in result.alerts if rule == "sheet_too_large"]
    assert len(alerts) == 2 and "上限 5 個" in alerts[0] and "行 2 までだけ" in alerts[0]
    assert result.changes == []  # 上限を超えた行（行 4）は読んでいない
    md = render.render(result, "markdown", "design-diff", "0")
    assert "上限 5 個" in md


def test_pair_by_duplicate_key_is_warned(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "API0001_旧"
    wb.create_sheet("API0001_新")
    wb.save(tmp_path / "b.xlsx")
    wb = Workbook()
    wb.active.title = "API0001_会員"
    wb.save(tmp_path / "a.xlsx")
    conf = {"sheets": [{"name": "API*", "pair_by": r"^(API\d+)", "items": []}]}
    result = excel_diff.run(conf, tmp_path / "b.xlsx", tmp_path / "a.xlsx")
    dup = [msg for rule, msg in result.alerts if rule == "duplicate_pair_key"]
    assert len(dup) == 1
    assert "b.xlsx" in dup[0] and "API0001" in dup[0] and "「API0001_旧」・「API0001_新」" in dup[0]
    assert ("sheet_renamed", "シート名", "API0001_旧", "API0001_会員") in [(c.category, c.field, c.before, c.after) for c in result.changes]


def test_single_cell_row_after_blank_sets_table_truncated(tmp_path: Path):
    """空行の後の 1 セルだけの行で表を打ち切ったときも table_truncated を出す（表の続きかもしれない）."""
    cells = table_cells(["項番", "項目名", "桁"], [[1, "id", 10], [2, "name", 50]])
    cells |= {"A5": "3", "A6": "補足"}
    before = book(tmp_path / "b.xlsx", cells)
    result = excel_diff.run(TABLE_CONF, before, before)
    trunc = [msg for rule, msg in result.alerts if rule == "table_truncated"]
    assert len(trunc) == 1 and "行 5" in trunc[0] and "A5「3」" in trunc[0]


def test_single_cell_row_that_is_a_configured_label_is_not_warned(tmp_path: Path):
    cells = table_cells(["項番", "項目名"], [[1, "id"]], start=2) | {"A1": "入力項目", "A4": "出力項目"}
    cells |= table_cells(["項番", "項目名"], [[1, "result"]], start=5)
    conf = {
        "sheets": [
            {
                "name": "IF*",
                "items": [
                    {"kind": "table", "name": "入力", "title": "入力項目", "header": "項番", "key": ["項番"]},
                    {"kind": "table", "name": "出力", "title": "出力項目", "header": "項番", "key": ["項番"]},
                ],
            }
        ]
    }
    p = book(tmp_path / "b.xlsx", cells)
    assert [r for r, _ in excel_diff.run(conf, p, p).alerts] == []


def test_global_yaml_paths_are_relative_to_the_config(tmp_path: Path, two_books, monkeypatch: pytest.MonkeyPatch):
    b, a1, _ = two_books
    conf_dir = _all_setup(tmp_path, [])
    g = {"output_dir": "../out", "pairs": [{"type": "if", "before": "../b.xlsx", "after": "../x/spec.xlsx"}]}
    (conf_dir / "global.yaml").write_text(yaml.safe_dump(g, allow_unicode=True), encoding="utf-8")
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)  # 実行場所に依らない
    rc, _, err = _run("excel", "--all", "--config-dir", str(conf_dir))
    assert rc == 0, err
    assert (tmp_path / "out" / "diff_spec.md").is_file()
    assert not (elsewhere / "out").exists()
    assert "相対" not in err


def test_global_yaml_legacy_cwd_relative_paths_still_work_with_warning(tmp_path: Path, two_books, monkeypatch: pytest.MonkeyPatch):
    b, a1, _ = two_books
    conf_dir = _all_setup(tmp_path, [])
    g = {"output_dir": "out", "pairs": [{"type": "if", "before": "b.xlsx", "after": "x/spec.xlsx"}]}  # 以前の書き方（実行場所から）
    (conf_dir / "global.yaml").write_text(yaml.safe_dump(g, allow_unicode=True), encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    rc, _, err = _run("excel", "--all", "--config-dir", str(conf_dir))
    assert rc == 0
    assert "実行場所からの相対として使いました" in err
    assert (conf_dir / "out" / "diff_spec.md").is_file()  # output_dir は global.yaml 基準（作るフォルダは常に存在しないため）


def test_python_dash_m_runs_the_cli():
    import subprocess  # nosec B404
    import sys

    root = Path(__file__).resolve().parent.parent
    proc = subprocess.run(  # nosec B603
        [sys.executable, "-m", "design_diff", "--version"], capture_output=True, text=True, encoding="utf-8", cwd=root, check=False
    )
    assert proc.returncode == 0 and "design-diff" in proc.stdout
    proc = subprocess.run(  # nosec B603
        [sys.executable, "-m", "design_diff", "excel", "--type", "api"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=root,
        check=False,
    )
    assert proc.returncode == 2  # 入力・使い方の誤り


def test_exit_codes_are_distinguished(tmp_path: Path, two_books, monkeypatch: pytest.MonkeyPatch):
    b, a1, _ = two_books
    conf_dir = _all_setup(tmp_path, [])
    args = ["--config-dir", str(conf_dir), "excel", "--type", "if", "--before", str(b)]
    assert _run(*args, "--after", str(a1))[0] == 0  # 差分ありでも既定は 0
    assert _run(*args, "--after", str(a1), "--fail-on-diff")[0] == 1  # 差分あり
    assert _run(*args, "--after", str(tmp_path / "none.xlsx"))[0] == 2  # 入力の誤り

    def boom(*_a, **_k):
        raise RuntimeError("想定外")

    monkeypatch.setattr(excel_diff, "run", boom)
    assert _run(*args, "--after", str(a1))[0] == 3  # 内部エラー
