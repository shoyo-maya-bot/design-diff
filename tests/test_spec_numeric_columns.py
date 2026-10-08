"""API 定義書の桁（max_length）: 表示形式に関係なくセルの値（数値）で読む. 「1,000」も数値. 読めなければ警告."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from design_diff.application import spec_diff
from design_diff.infrastructure import excel_spec_reader

ROOT = Path(__file__).resolve().parent.parent
CONF = yaml.safe_load((ROOT / "config/spec.yaml").read_text(encoding="utf-8"))
OAS = """openapi: 3.0.3
info: {title: t, version: '1'}
paths:
  /items:
    get:
      x-api-id: API0001
      responses:
        '200':
          content:
            application/json:
              schema:
                type: object
                properties:
                  name: {type: string, maxLength: 1000}
                  code: {type: string, maxLength: 5}
"""


def _book(path: Path, lengths: list[tuple[str, object, str | None]]) -> Path:
    """出力項目の表に（項目名, 桁の値, 桁の表示形式）を並べた API 定義書."""
    wb = Workbook()
    ws = wb.active
    ws.title = "API0001"
    ws["A1"], ws["B1"] = "API ID", "API0001"
    ws["A2"], ws["B2"] = "メソッド", "GET"
    ws["A3"], ws["B3"] = "パス", "/items"
    ws["A5"] = "出力項目"
    for c, h in enumerate(["項目名", "型", "必須", "桁"], start=1):
        ws.cell(6, c, h)
    for r, (name, length, fmt) in enumerate(lengths, start=7):
        ws.cell(r, 1, name)
        ws.cell(r, 2, "文字列")
        cell = ws.cell(r, 4, length)
        if fmt:
            cell.number_format = fmt
    wb.save(path)
    return path


@pytest.mark.parametrize(
    ("value", "fmt"),
    [
        (1000, "#,##0"),  # 表示は「1,000」
        (1000, "0000000"),  # 表示は「0001000」
        (1000, '0"桁"'),  # 表示は「1000桁」
        ("1,000", None),  # 文字列で入力した「1,000」
        ("１０００", None),  # 全角
    ],
)
def test_length_is_read_as_number_regardless_of_format(tmp_path: Path, value: object, fmt: str | None):
    spec, _, _, alerts = excel_spec_reader.read_with_alerts(_book(tmp_path / "a.xlsx", [("name", value, fmt)]), CONF)
    assert spec["API0001"]["response"]["name"]["max_length"] == 1000
    assert not [a for a in alerts if a[0] == "not_a_number"]


def test_formatted_length_is_compared(tmp_path: Path):
    oas = tmp_path / "openapi.yaml"
    oas.write_text(OAS, encoding="utf-8")
    book = _book(tmp_path / "a.xlsx", [("name", 1000, "#,##0"), ("code", 1500, "#,##0")])
    result = spec_diff.run(CONF, book, oas)
    mismatches = [(c.field, c.before, c.after) for c in result.changes if c.category == "length_mismatch"]
    assert mismatches == [("code", "1500", "5")]  # name は同じ（1,000 = 1000）、code は違う（以前は比べていなかった）


def test_unreadable_length_is_warned(tmp_path: Path):
    spec, _, _, alerts = excel_spec_reader.read_with_alerts(_book(tmp_path / "a.xlsx", [("name", "可変", None)]), CONF)
    assert spec["API0001"]["response"]["name"]["max_length"] is None
    warned = [w for rule, w in alerts if rule == "not_a_number"]
    assert len(warned) == 1 and "name" in warned[0] and "可変" in warned[0]
