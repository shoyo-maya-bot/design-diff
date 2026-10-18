"""第 6 回: block のラベル直下の空行・型と位置の大小文字・head/options/trace・2xx の選び方."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from test_robustness import book, changes_of

from design_diff.domain import spec_diff
from design_diff.infrastructure import excel_reader, oas_reader

BLOCK_CONF = {"sheets": [{"name": "IF*", "items": [{"kind": "block", "name": "処理概要", "label": "処理概要"}]}]}


# --- block: ラベルの直下が空行 ----------------------------------------------------------------


def test_block_skips_blank_row_right_under_label(tmp_path: Path):
    base = {"A1": "処理概要", "A3": "1 行目", "A4": "2 行目"}
    wb = Workbook()
    ws = wb.active
    for ref, v in base.items():
        ws[ref] = v
    ctx = excel_reader.ReadContext()
    assert excel_reader.read_block(ws, "処理概要", ctx=ctx) == "1 行目\n2 行目"  # 以前は空文字
    assert any("直下の空行（行 2）" in n for n in ctx.notes)


def test_block_under_blank_row_changes_are_detected(tmp_path: Path):
    base = {"A1": "処理概要", "A3": "受け付けたデータを登録する"}
    before = book(tmp_path / "b.xlsx", base)
    after = book(tmp_path / "a.xlsx", base | {"A3": "受け付けたデータを削除する"})
    (cat, _, b, a), *rest = changes_of(BLOCK_CONF, before, after)
    assert (cat, b, a, rest) == ("value_changed", "受け付けたデータを登録する", "受け付けたデータを削除する", [])


def test_block_blank_then_table_is_still_empty(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A3"], ws["B3"] = "処理概要", "項番", "項目名"  # 空行の後に表（別の列にも値）
    assert excel_reader.read_block(ws, "処理概要") == ""


# --- 型・位置の大小文字 ------------------------------------------------------------------------


def _api(fields: dict) -> dict:
    return {"API0001": {"method": "GET", "path": "/x", "request": fields, "response": {}, "statuses": []}}


def test_type_and_location_compare_ignoring_case():
    doc = _api({"id": {"type": "String", "location": "Query", "required": True}})
    oas = _api({"id": {"type": "string", "location": "query", "required": True}})
    assert spec_diff.compare(doc, oas) == []
    oas2 = _api({"id": {"type": "integer", "location": "path", "required": True}})
    cats = sorted(c.category for c in spec_diff.compare(doc, oas2))
    assert cats == ["location_mismatch", "type_mismatch"]


# --- head / options / trace -------------------------------------------------------------------

OAS_METHODS = """openapi: 3.0.3
info: {title: t, version: '1'}
paths:
  /items:
    head: {operationId: API0010, responses: {'200': {}}}
    options: {operationId: API0011, responses: {'204': {}}}
    trace: {operationId: API0012, responses: {'200': {}}}
"""


def test_head_options_trace_are_read(tmp_path: Path):
    p = tmp_path / "o.yaml"
    p.write_text(OAS_METHODS, encoding="utf-8")
    spec = oas_reader.read(p)
    assert {k: v["method"] for k, v in spec.items()} == {"API0010": "HEAD", "API0011": "OPTIONS", "API0012": "TRACE"}


# --- 2xx の選び方 -----------------------------------------------------------------------------

OAS_2XX = """openapi: 3.0.3
info: {{title: t, version: '1'}}
paths:
  /jobs:
    post:
      x-api-id: API0001
      responses:
{responses}
"""

BODY_200 = """        '200':
          content:
            application/json:
              schema: {type: object, properties: {id: {type: integer}}}
"""


def _read_2xx(tmp_path: Path, responses: str) -> tuple[dict, list[str]]:
    p = tmp_path / "o.yaml"
    p.write_text(OAS_2XX.format(responses=responses), encoding="utf-8")
    spec, notes = oas_reader.read_with_notes(p)
    return spec["API0001"], notes


def test_202_without_body_before_200_uses_200_schema(tmp_path: Path):
    api, notes = _read_2xx(tmp_path, "        '202': {description: accepted}\n" + BODY_200)
    assert api["response"]["id"]["type"] == "integer"  # 以前は 202 を見て項目なし
    assert api["statuses"] == ["200", "202"]
    assert not any("成功応答が複数" in n for n in notes)


def test_unquoted_numeric_codes_and_multiple_bodies_note(tmp_path: Path):
    responses = (
        "        201:\n          content:\n            application/json:\n"
        "              schema: {type: object, properties: {name: {type: string}}}\n" + BODY_200.replace("'200'", "200")
    )
    api, notes = _read_2xx(tmp_path, responses)
    assert set(api["response"]) == {"id"}  # 200 を優先
    assert any("200 で比べました（他: 201）" in n for n in notes)
