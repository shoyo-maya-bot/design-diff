"""第 3 回の指摘の回帰テスト: キー列が空の行、OpenAPI の壊れた入力、同じ API ID のシート."""

from __future__ import annotations

import contextlib
import io
from pathlib import Path

import pytest
from devtools_common.cli import run_main
from openpyxl import Workbook
from test_robustness import TABLE_CONF, book, table_cells

from design_diff.application import excel_diff, spec_diff
from design_diff.cli import main as cli
from design_diff.domain.excel_diff import blank_key_rows

ROOT = Path(__file__).resolve().parent.parent
SPEC_CONF = {
    "sheets": "API*",
    "api_id": {"label": "API ID"},
    "method": {"label": "メソッド"},
    "path": {"label": "パス"},
}


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue() + err.getvalue()


# --- 1. キー列が空の行 --------------------------------------------------------------------------


def test_blank_key_rows():
    rows = [{"項番": "1"}, {"項番": ""}, {"項番": " "}, {"項番": "2"}]
    assert blank_key_rows(["項番"], rows) == [1, 2]
    assert blank_key_rows(["a", "b"], [{"a": "", "b": "x"}, {"a": "", "b": ""}]) == [1]


def test_blank_key_rows_are_warned_with_count_and_row_numbers(tmp_path: Path):
    header = ["項番", "名前", "型"]
    before = book(tmp_path / "b.xlsx", table_cells(header, [[1, "a", "数値"], [None, "b", "文字列"], [2, "c", "数値"]]))
    after = book(
        tmp_path / "a.xlsx", table_cells(header, [[1, "a", "数値"], [None, "b2", "文字列"], [None, "x", "日付"], [2, "c", "数値"]])
    )
    result = excel_diff.run(TABLE_CONF, before, after)
    alerts = [w for rule, w in result.alerts if rule == "blank_key"]
    assert len(alerts) == 2, result.alerts
    assert "変更前" in alerts[0] and "1 行" in alerts[0] and "行 3）" in alerts[0]
    assert "変更後" in alerts[1] and "2 行" in alerts[1] and "行 3、4）" in alerts[1]
    # 空キーでない行に差分が無ければ、差分は 0 件（空キーの行は比べていないことを警告で知らせる）
    assert [c for c in result.changes if c.category not in ("outside_changed",)] == []


def test_no_blank_key_no_alert(tmp_path: Path):
    header = ["項番", "名前"]
    b = book(tmp_path / "b.xlsx", table_cells(header, [[1, "a"]]))
    assert not [r for r, _ in excel_diff.run(TABLE_CONF, b, b).alerts if r == "blank_key"]


# --- 2. OpenAPI の壊れた入力は終了コード 2 --------------------------------------------------------


@pytest.mark.parametrize(
    "content",
    [
        "openapi: 3.0.0\npaths: {\n  /x: [\n",  # YAML の構文の誤り
        "openapi: 3.0.0\npaths:\n  /x: [1, 2]\n",  # パスの中身が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get: 1\n",  # 操作が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get:\n      parameters: [1]\n",  # パラメーターが辞書でない
        "openapi: 3.0.0\npaths: [1]\n",  # paths が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get:\n      responses: {'200': {content: {application/json: {schema: {$ref: '#/nope'}}}}}\n",
    ],
)
def test_broken_openapi_is_input_error(tmp_path: Path, content: str):
    oas = tmp_path / "openapi.yaml"
    oas.write_text(content, encoding="utf-8")
    rc, out = run("spec", "--excel", str(ROOT / "samples/spec/API定義書.xlsx"), "--oas", str(oas))
    assert rc == 2, out
    assert "Traceback" not in out


def test_broken_openapi_raises_value_error(tmp_path: Path):
    from design_diff.infrastructure import oas_reader

    oas = tmp_path / "o.yaml"
    oas.write_text("paths: {\n", encoding="utf-8")
    with pytest.raises(ValueError, match="OpenAPI"):
        oas_reader.read_with_notes(oas)


# --- 3. 同じ API ID のシート ------------------------------------------------------------------------


def test_duplicate_api_id_uses_first_sheet_and_warns(tmp_path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    for title, method in (("API1", "GET"), ("API2", "POST")):
        ws = wb.create_sheet(title)
        ws["A1"], ws["B1"] = "API ID", "API0001"
        ws["A2"], ws["B2"] = "メソッド", method
        ws["A3"], ws["B3"] = "パス", "/x"
    path = tmp_path / "spec.xlsx"
    wb.save(path)
    oas = tmp_path / "o.yaml"
    oas.write_text("openapi: 3.0.0\npaths:\n  /x:\n    get:\n      operationId: API0001\n      responses: {'200': {}}\n", encoding="utf-8")
    result = spec_diff.run(SPEC_CONF, path, oas)
    alerts = [w for rule, w in result.alerts if rule == "duplicate_api_id"]
    assert len(alerts) == 1 and "「API1」" in alerts[0] and "「API2」" in alerts[0]
    assert not [c for c in result.changes if c.category == "method_mismatch"]  # 最初のシート（GET）を使う
