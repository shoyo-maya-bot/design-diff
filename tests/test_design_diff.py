"""design-diff のテスト. サンプル生成スクリプトで一時ディレクトリに設計書を作り、差分の種類と終了コードを確かめる."""

from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import shutil
from pathlib import Path

import pytest
import yaml
from devtools_common.cli import run_main
from openpyxl import load_workbook

from design_diff.application import excel_diff
from design_diff.cli import main as cli
from design_diff.domain.excel_diff import compare_table, compare_value
from design_diff.domain.spec_diff import compare

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config"

_spec = importlib.util.spec_from_file_location("make_samples", ROOT / "samples/make_samples.py")
assert _spec and _spec.loader
make_samples = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(make_samples)


@pytest.fixture()
def samples(tmp_path: Path) -> Path:
    make_samples.generate(tmp_path)
    return tmp_path


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, ["--config-dir", str(CONFIG), *args])
    return rc, out.getvalue() + err.getvalue()


def categories(out: str) -> dict[str, list[dict]]:
    data = json.loads(out)
    groups: dict[str, list[dict]] = {}
    for f in data["findings"]:
        groups.setdefault(f["rule_id"], []).append(f)
    return groups


# --- ドメイン（抽出済みの値の比較） ---------------------------------------------------------


def test_compare_value_kinds():
    assert compare_value("s", "x", "a", "a") == []
    assert compare_value("s", "x", "a", "b")[0].category == "value_changed"
    assert compare_value("s", "x", None, "b")[0].category == "label_added"
    assert compare_value("s", "x", "a", None)[0].category == "label_removed"
    assert compare_value("s", "x", None, None)[0].category == "label_not_found"


def test_compare_value_ignores_surrounding_spaces_and_float_ints():
    from design_diff.domain.model import text

    assert compare_value("s", "x", text(" 10 "), text(10.0)) == []


def test_compare_table_rows_cells_columns():
    before = (["No", "名前", "桁"], [{"No": "1", "名前": "a", "桁": "10"}, {"No": "2", "名前": "b", "桁": "5"}])
    after = (
        ["No", "名前", "桁", "備考"],
        [{"No": "1", "名前": "a", "桁": "12", "備考": ""}, {"No": "3", "名前": "c", "桁": "1", "備考": ""}],
    )
    cats = [c.category for c in compare_table("t", ["No"], before, after)]
    assert sorted(cats) == ["cell_changed", "column_added", "row_added", "row_removed"]


def test_compare_table_missing_key_and_duplicate():
    t = (["No"], [{"No": "1"}, {"No": "1"}])
    assert [c.category for c in compare_table("t", ["ID"], t, t)] == ["key_column_missing"]
    assert "duplicate_key" in [c.category for c in compare_table("t", ["No"], t, t)]


def test_spec_compare_path_params_are_normalized():
    doc = {"A": {"method": "get", "path": "/x/{id}/", "request": {}, "response": {}, "statuses": []}}
    oas = {"A": {"method": "GET", "path": "/x/{memberId}", "request": {}, "response": {}, "statuses": ["200"]}}
    assert compare(doc, oas) == []  # 設計書に応答一覧が無ければステータスは比べない


# --- Excel の新旧差分 ----------------------------------------------------------------------


def test_api_excel_diff(samples: Path):
    rc, out = run(
        "excel", "--type", "api", "--before", str(samples / "before/API0001_会員照会.xlsx"),
        "--after", str(samples / "after/API0001_会員照会.xlsx"), "-f", "json",
    )  # fmt: skip
    assert rc == 0
    g = categories(out)
    assert set(g) == {"sheet_renamed", "value_changed", "row_added", "row_removed", "cell_changed"}
    assert {f["field"] for f in g["value_changed"]} == {"API 名", "処理概要"}
    assert {f["field"] for f in g["row_added"]} == {"includeHistory", "createdAt", "400", "1.1"}
    assert [f["field"] for f in g["row_removed"]] == ["faxNumber"]
    assert (g["cell_changed"][0]["field"], g["cell_changed"][0]["expected"], g["cell_changed"][0]["actual"]) == ("name の 桁", "50", "60")


def test_batch_excel_diff_markdown(samples: Path):
    rc, out = run(
        "excel", "--type", "batch", "--before", str(samples / "before/JOB0001_注文取込.xlsx"),
        "--after", str(samples / "after/JOB0001_注文取込.xlsx"),
    )  # fmt: skip
    assert rc == 0
    assert "| 概要 | 起動条件 | 毎日 02:00（スケジューラ） | 毎日 01:30（スケジューラ） |" in out
    assert "警告終了" in out


def test_same_file_has_no_diff_and_fail_on_diff(samples: Path):
    f = str(samples / "before/JOB0001_注文取込.xlsx")
    rc, out = run("excel", "--type", "batch", "--before", f, "--after", f, "--fail-on-diff")
    assert (rc, "差分はありません" in out) == (0, True)
    rc, _ = run("excel", "--type", "batch", "--before", f, "--after", str(samples / "after/JOB0001_注文取込.xlsx"), "--fail-on-diff")
    assert rc == 1


def test_label_not_found_is_config_problem(samples: Path, tmp_path: Path):
    conf_dir = tmp_path / "conf"
    conf_dir.mkdir()
    (conf_dir / "x.yaml").write_text(
        yaml.safe_dump({"sheets": [{"name": "概要", "items": [{"kind": "label", "name": "担当", "label": "担当者"}]}]}, allow_unicode=True),
        encoding="utf-8",
    )
    f = str(samples / "before/JOB0001_注文取込.xlsx")
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, ["--config-dir", str(conf_dir), "excel", "--type", "x", "--before", f, "--after", f])
    assert rc == 1  # 差分が無くても、設定の見直しが必要なら 1
    assert "設定の見直しが必要" in out.getvalue()


def test_merged_cells_are_read(samples: Path):
    p = samples / "after/JOB0001_注文取込.xlsx"
    wb = load_workbook(p)
    ws = wb["概要"]
    ws["C5"] = None
    ws.merge_cells("B5:C5")
    wb.save(p)
    rc, out = run("excel", "--type", "batch", "--before", str(samples / "before/JOB0001_注文取込.xlsx"), "--after", str(p))
    assert rc == 0 and "毎日 01:30" in out


def test_all_pairs(samples: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    conf_dir = tmp_path / "conf"
    shutil.copytree(CONFIG, conf_dir)
    g = yaml.safe_load((conf_dir / "global.yaml").read_text(encoding="utf-8"))
    g["output_dir"] = str(tmp_path / "out")
    for pair in g["pairs"]:
        for side in ("before", "after"):
            pair[side] = str(samples / Path(pair[side]).relative_to("../samples"))
    (conf_dir / "global.yaml").write_text(yaml.safe_dump(g, allow_unicode=True), encoding="utf-8")
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        rc = run_main(cli.main, ["--config-dir", str(conf_dir), "excel", "--all"])
    assert rc == 0
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == ["diff_API0001_会員照会.md", "diff_JOB0001_注文取込.md"]


@pytest.mark.parametrize(
    ("item", "message"),
    [
        ({"kind": "cell", "name": "x"}, "kind"),
        ({"kind": "label", "name": "x"}, "label"),
        ({"kind": "table", "name": "x", "header": "No"}, "header と key"),
    ],
)
def test_validate_config(item: dict, message: str):
    with pytest.raises(ValueError, match=message):
        excel_diff.validate_config({"sheets": [{"name": "*", "items": [item]}]}, "t.yaml")


def test_validate_pair_by_needs_group():
    with pytest.raises(ValueError, match="グループ"):
        excel_diff.validate_config({"sheets": [{"name": "*", "pair_by": "API\\d+", "items": []}]}, "t.yaml")


def test_shipped_configs_are_valid():
    for name in ("api", "batch", "if"):
        excel_diff.validate_config(yaml.safe_load((CONFIG / f"{name}.yaml").read_text(encoding="utf-8")), name)


def test_usage_errors(samples: Path):
    assert run("excel", "--type", "api")[0] == 2
    assert run("excel", "--type", "nothing", "--before", "a", "--after", "b")[0] == 2
    assert run("excel", "--type", "api", "--before", "none.xlsx", "--after", "none.xlsx")[0] == 2


# --- API 定義書と OpenAPI --------------------------------------------------------------------


def test_spec_diff(samples: Path):
    rc, out = run("spec", "--excel", str(samples / "spec/API定義書.xlsx"), "--oas", str(samples / "spec/openapi.yaml"), "-f", "json")
    assert rc == 0
    g = categories(out)
    assert set(g) - {"note"} == {
        "api_missing_in_doc", "length_mismatch", "status_missing_in_oas", "field_missing_in_doc", "required_mismatch",
    }  # fmt: skip
    assert g["api_missing_in_doc"][0]["location"] == "API0003"
    assert (g["length_mismatch"][0]["field"], g["length_mismatch"][0]["expected"], g["length_mismatch"][0]["actual"]) == (
        "name",
        "60",
        "50",
    )
    assert (g["required_mismatch"][0]["location"], g["required_mismatch"][0]["field"]) == ("API0002 request", "email")
    assert g["field_missing_in_doc"][0]["field"] == "phone"


def test_spec_diff_detects_type_location_and_method(samples: Path):
    p = samples / "spec/openapi.yaml"
    text = p.read_text(encoding="utf-8")
    text = text.replace("        - name: includeHistory\n          in: query\n          schema: {type: boolean}",
                        "        - name: includeHistory\n          in: header\n          schema: {type: string}")  # fmt: skip
    text = text.replace("  /api/v1/members:\n    post:", "  /api/v1/members:\n    put:")
    p.write_text(text, encoding="utf-8")
    rc, out = run("spec", "--excel", str(samples / "spec/API定義書.xlsx"), "--oas", str(p), "-f", "json")
    g = categories(out)
    assert {"type_mismatch", "location_mismatch", "method_mismatch"} <= set(g)


def test_spec_bad_openapi_is_usage_error(samples: Path):
    bad = samples / "bad.yaml"
    bad.write_text("openapi: 3.0.3\n", encoding="utf-8")
    assert run("spec", "--excel", str(samples / "spec/API定義書.xlsx"), "--oas", str(bad))[0] == 2


# --- 2 段の見出し（結合セルのグループ見出し + 子見出し） ------------------------------------------


def _two_row_header_book(path: Path, rows: list[list], *, merge_vertical: bool = True) -> Path:
    """1 段目: 項目名（縦結合）・型情報（横結合）・備考（縦結合）、2 段目: 型・桁 の表."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "IF0001"
    ws["A1"] = "項目名"
    ws["B1"] = "型情報"
    ws["D1"] = "備考"
    ws["B2"] = "型"
    ws["C2"] = "桁"
    ws.merge_cells("B1:C1")
    if merge_vertical:
        ws.merge_cells("A1:A2")
        ws.merge_cells("D1:D2")
    for i, row in enumerate(rows, start=3):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    return path


ROWS = [["id", "数値", 10, ""], ["name", "文字列", 50, "氏名"]]


@pytest.mark.parametrize("merge_vertical", [True, False])
def test_two_row_header_is_combined(tmp_path: Path, merge_vertical: bool):
    from design_diff.infrastructure import excel_reader

    ws = excel_reader.open_book(_two_row_header_book(tmp_path / "a.xlsx", ROWS, merge_vertical=merge_vertical)).active
    cols, rows = excel_reader.read_table(ws, "項目名")
    assert cols == ["項目名", "型情報/型", "型情報/桁", "備考"]
    assert rows == [
        {"項目名": "id", "型情報/型": "数値", "型情報/桁": "10", "備考": ""},
        {"項目名": "name", "型情報/型": "文字列", "型情報/桁": "50", "備考": "氏名"},
    ]


def test_two_row_header_diff_reports_the_sub_column(tmp_path: Path):
    before = _two_row_header_book(tmp_path / "b.xlsx", ROWS)
    after = _two_row_header_book(tmp_path / "a.xlsx", [["id", "数値", 10, ""], ["name", "文字列", 60, "氏名"]])
    conf = {"sheets": [{"name": "IF*", "items": [{"kind": "table", "name": "レイアウト", "header": "項目名", "key": ["項目名"]}]}]}
    changes = excel_diff.run(conf, before, after).changes
    assert [(c.category, c.field, c.before, c.after) for c in changes] == [("cell_changed", "name の 型情報/桁", "50", "60")]


def test_header_rows_option(tmp_path: Path):
    """header_rows で段数を指定できる. 1 にすると従来どおり 1 段目だけを見出しにする."""
    from design_diff.infrastructure import excel_reader

    ws = excel_reader.open_book(_two_row_header_book(tmp_path / "a.xlsx", ROWS, merge_vertical=False)).active
    cols, rows = excel_reader.read_table(ws, "項目名", header_rows=1)
    assert cols == ["項目名", "型情報", "備考"]
    cols, _ = excel_reader.read_table(ws, "項目名", header_rows=2)
    assert cols == ["項目名", "型情報/型", "型情報/桁", "備考"]


def test_single_row_header_with_merged_column_is_unchanged(tmp_path: Path):
    """見出しが横結合でも、下の行が結合のままのデータなら 1 段の見出しとして読む."""
    from openpyxl import Workbook

    from design_diff.infrastructure import excel_reader

    wb = Workbook()
    ws = wb.active
    ws.append(["No", "説明", None])
    ws.append([1, "最初", None])
    ws.append([2, "次", None])
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    wb.save(tmp_path / "s.xlsx")
    cols, rows = excel_reader.read_table(excel_reader.open_book(tmp_path / "s.xlsx").active, "No")
    assert cols == ["No", "説明"]
    assert rows == [{"No": "1", "説明": "最初"}, {"No": "2", "説明": "次"}]


def test_validate_header_rows():
    item = {"kind": "table", "name": "x", "header": "No", "key": ["No"], "header_rows": 0}
    with pytest.raises(ValueError, match="header_rows"):
        excel_diff.validate_config({"sheets": [{"name": "*", "items": [item]}]}, "t.yaml")


# --- 表示形式どおりの値で比べる ---------------------------------------------------------------------


def _code_book(path: Path, code: object, fmt: str = "General", amount: object = 1500, amount_fmt: str = "#,##0") -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "IF0001"
    ws.append(["コード", "金額"])
    ws.append([code, amount])
    ws["A2"].number_format = fmt
    ws["B2"].number_format = amount_fmt
    wb.save(path)
    return path


CODE_CONF = {"sheets": [{"name": "IF*", "items": [{"kind": "table", "name": "コード表", "header": "コード", "key": ["金額"]}]}]}


def test_text_and_zero_padded_number_are_equal(tmp_path: Path):
    """文字列の「001」と、表示形式 000 の数値 1 は Excel の画面で同じに見えるので差分にしない."""
    before = _code_book(tmp_path / "b.xlsx", "001", "@")
    after = _code_book(tmp_path / "a.xlsx", 1, "000", amount="1,500", amount_fmt="@")
    assert excel_diff.run(CODE_CONF, before, after).changes == []


def test_display_format_change_is_reported(tmp_path: Path):
    """値が同じでも、表示が変わる表示形式の変更（000 → 0）は差分として出す. 値は表示どおりの文字列."""
    before = _code_book(tmp_path / "b.xlsx", 1, "000")
    after = _code_book(tmp_path / "a.xlsx", 1, "0")
    changes = excel_diff.run(CODE_CONF, before, after).changes
    assert [(c.category, c.field, c.before, c.after) for c in changes] == [("cell_changed", "1,500 の コード", "001", "1")]
