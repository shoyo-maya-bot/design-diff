"""表の行のうち、表の列の範囲の外のセル（表の右のメモ・承認欄など）は「設定に無いセル」として比べる."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from test_robustness import TABLE_CONF, book, changes_of, table_cells

from design_diff.infrastructure import excel_reader

HEADER = ["項番", "項目名", "型"]
ROWS = [[1, "id", "数値"], [2, "name", "文字列"]]


def test_cell_right_of_table_row_is_compared_as_outside(tmp_path: Path):
    base = table_cells(HEADER, ROWS)
    before = book(tmp_path / "b.xlsx", base | {"F2": "確認済み"})
    after = book(tmp_path / "a.xlsx", base | {"F2": "差し戻し"})
    got = changes_of(TABLE_CONF, before, after, outside="all")
    assert got == [("outside_cell_changed", "F2", "確認済み", "差し戻し")]


def test_cell_inside_table_columns_is_compared_as_table(tmp_path: Path):
    before = book(tmp_path / "b.xlsx", table_cells(HEADER, ROWS))
    after = book(tmp_path / "a.xlsx", table_cells(HEADER, [[1, "id", "数値"], [2, "name", "日付"]]))
    assert [c[0] for c in changes_of(TABLE_CONF, before, after, outside="all")] == ["cell_changed"]


def test_merged_header_extends_table_range(tmp_path: Path):
    """横に結合した見出しの右端の列までは表の中（データの結合も表の値として読む）."""
    wb = Workbook()
    ws = wb.active
    ws.title = "IF0001"
    for ref, v in {"A1": "項番", "B1": "説明", "A2": 1, "B2": "x", "C2": "表の中", "E2": "表の外"}.items():
        ws[ref] = v
    ws.merge_cells("B1:C1")
    wb.save(tmp_path / "a.xlsx")
    ws = excel_reader.open_book(tmp_path / "a.xlsx").active
    ctx = excel_reader.ReadContext()
    excel_reader.read_table(ws, "項番", header_rows=1, ctx=ctx)
    assert (2, 3) in ctx.covered and (2, 5) not in ctx.covered
    assert [r for r, _ in excel_reader.outside_rows(ws, ctx.covered)] == [2]
