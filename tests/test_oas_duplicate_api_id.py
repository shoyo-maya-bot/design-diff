"""OpenAPI 側で同じ x-api-id の操作が 2 つ以上あると、最初の操作を使い、警告 duplicate_api_id を出す（黙って上書きしない）."""

from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path

from devtools_common.cli import run_main
from openpyxl import Workbook

from design_diff.application import spec_diff
from design_diff.cli import main as cli
from design_diff.infrastructure import oas_reader

SPEC_CONF = {"sheets": "API*", "api_id": {"label": "API ID"}, "method": {"label": "メソッド"}, "path": {"label": "パス"}}

OAS = """openapi: 3.0.0
paths:
  /members/{id}:
    get:
      x-api-id: API0001
      responses: {'200': {}}
  /orders:
    post:
      x-api-id: API0001
      responses: {'201': {}}
"""


def _excel(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "API1"
    ws["A1"], ws["B1"] = "API ID", "API0001"
    ws["A2"], ws["B2"] = "メソッド", "GET"
    ws["A3"], ws["B3"] = "パス", "/members/{id}"
    path = tmp_path / "spec.xlsx"
    wb.save(path)
    return path


def test_reader_keeps_first_operation_and_reports_duplicate(tmp_path: Path):
    oas = tmp_path / "o.yaml"
    oas.write_text(OAS, encoding="utf-8")
    spec, _notes, alerts = oas_reader.read_with_alerts(oas)
    assert spec["API0001"]["method"] == "GET" and spec["API0001"]["path"] == "/members/{id}"
    dup = [w for rule, w in alerts if rule == "duplicate_api_id"]
    assert len(dup) == 1
    assert "GET /members/{id}" in dup[0] and "POST /orders" in dup[0]


def test_spec_diff_shows_duplicate_alert_and_compares_first(tmp_path: Path):
    oas = tmp_path / "o.yaml"
    oas.write_text(OAS, encoding="utf-8")
    result = spec_diff.run(SPEC_CONF, _excel(tmp_path), oas)
    assert [r for r, _ in result.alerts].count("duplicate_api_id") == 1
    assert not [c for c in result.changes if c.category in ("method_mismatch", "path_mismatch")]


def test_cli_json_has_duplicate_alert(tmp_path: Path):
    oas = tmp_path / "o.yaml"
    oas.write_text(OAS, encoding="utf-8")
    out = io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
        rc = run_main(cli.main, ["spec", "--excel", str(_excel(tmp_path)), "--oas", str(oas), "-f", "json"])
    assert rc in (0, 1)
    assert "duplicate_api_id" in out.getvalue()
    json.loads(out.getvalue())


def test_unique_ids_have_no_alert(tmp_path: Path):
    oas = tmp_path / "o.yaml"
    unique = OAS.replace("x-api-id: API0001\n      responses: {'201'", "x-api-id: API0002\n      responses: {'201'")
    assert unique != OAS
    oas.write_text(unique, encoding="utf-8")
    _spec, _notes, alerts = oas_reader.read_with_alerts(oas)
    assert alerts == []
