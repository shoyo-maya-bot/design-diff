"""計算結果が保存されていない数式セルの扱いのテスト（openpyxl で作った Excel は計算結果を持たない）."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from design_diff.application import excel_diff, render
from design_diff.infrastructure import excel_reader

CONF = {
    "sheets": [
        {
            "name": "*",
            "items": [
                {"kind": "label", "name": "合計", "label": "合計"},
                {"kind": "table", "name": "明細", "header": "項目", "key": ["項目"]},
            ],
        }
    ]
}


def _book(path: Path, total_formula: str, price_formula: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"] = "合計"
    ws["B1"] = total_formula
    ws["A3"] = "項目"
    ws["B3"] = "金額"
    ws["A4"] = "単価"
    ws["B4"] = 100
    ws["A5"] = "税込"
    ws["B5"] = price_formula
    wb.save(path)
    return path


def test_load_book_uses_formula_text_when_cache_is_missing(tmp_path: Path):
    p = _book(tmp_path / "a.xlsx", "=SUM(B4:B5)", "=B4*1.1")
    wb, warnings = excel_reader.load_book(p)
    assert wb["S1"]["B1"].value == "=SUM(B4:B5)"
    assert wb["S1"]["B4"].value == 100
    assert len(warnings) == 1
    assert "a.xlsx" in warnings[0] and "2 個" in warnings[0] and "S1!B1" in warnings[0] and "保存" in warnings[0]


def test_book_without_formulas_has_no_warning(tmp_path: Path):
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(tmp_path / "plain.xlsx")
    assert excel_reader.load_book(tmp_path / "plain.xlsx")[1] == []


def test_excel_diff_compares_formula_text_and_warns(tmp_path: Path):
    b = _book(tmp_path / "before.xlsx", "=SUM(B4:B5)", "=B4*1.1")
    a = _book(tmp_path / "after.xlsx", "=SUM(B4:B5)", "=B4*1.08")
    result = excel_diff.run(CONF, b, a)
    assert [(c.category, c.field, c.before, c.after) for c in result.changes] == [("cell_changed", "税込 の 金額", "=B4*1.1", "=B4*1.08")]
    assert len(result.warnings) == 2
    data = json.loads(render.render(result, "json", "design-diff", "0"))
    assert [f["severity"] for f in data["findings"] if f["rule_id"] == "formula_cache_missing"] == ["warning", "warning"]
    md = render.markdown(result)
    assert "注意" in md and "Excel で開いて保存" in md


def test_same_formulas_have_no_diff(tmp_path: Path):
    b = _book(tmp_path / "before.xlsx", "=SUM(B4:B5)", "=B4*1.1")
    result = excel_diff.run(CONF, b, b)
    assert result.changes == []
    assert "差分はありません" in render.markdown(result)
    assert result.warnings
