"""spec 比較の項目は名前と受け渡し位置（path / query / body など）で識別する.

path パラメーターと body に同じ名前の項目があっても 1 つにまとめない（以前は後の方で上書きされ、片方を比べていなかった）.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from design_diff.application import spec_diff as app_spec_diff
from design_diff.domain import spec_diff
from design_diff.infrastructure import excel_spec_reader, oas_reader

CONF = {
    "sheets": "API*",
    "api_id": {"label": "API ID"},
    "method": {"label": "メソッド"},
    "path": {"label": "パス"},
    "request": {
        "title": "入力項目",
        "columns": {"name": "項目名", "type": "型", "required": "必須", "max_length": "桁", "location": "位置"},
    },
    "type_map": {"文字列": "string", "数値": "integer"},
}

OAS = """openapi: 3.0.0
paths:
  /members/{id}:
    put:
      x-api-id: API0001
      parameters:
        - {name: id, in: path, required: true, schema: {type: string, maxLength: 10}}
      requestBody:
        content:
          application/json:
            schema:
              type: object
              required: [id]
              properties:
                id: {type: integer}
                name: {type: string, maxLength: 20}
      responses: {'200': {}}
"""


def _excel(tmp_path: Path, rows: list[tuple]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "API1"
    ws["A1"], ws["B1"] = "API ID", "API0001"
    ws["A2"], ws["B2"] = "メソッド", "PUT"
    ws["A3"], ws["B3"] = "パス", "/members/{id}"
    ws["A5"] = "入力項目"
    ws.append(["項目名", "型", "必須", "桁", "位置"])
    for r in rows:
        ws.append(list(r))
    path = tmp_path / "spec.xlsx"
    wb.save(path)
    return path


def _oas(tmp_path: Path) -> Path:
    p = tmp_path / "o.yaml"
    p.write_text(OAS, encoding="utf-8")
    return p


def test_oas_reader_keeps_path_parameter_and_body_field_with_same_name(tmp_path: Path):
    spec, _ = oas_reader.read_with_notes(_oas(tmp_path))
    ids = [f for k, f in spec["API0001"]["request"].items() if f.get("name", k) == "id"]
    assert sorted((f["location"], f["type"]) for f in ids) == [("body", "integer"), ("path", "string")]


def test_excel_reader_keeps_rows_with_same_name_and_different_location(tmp_path: Path):
    path = _excel(tmp_path, [("id", "文字列", "○", 10, "path"), ("id", "数値", "○", None, "body")])
    spec, _, _ = excel_spec_reader.read(path, CONF)
    ids = [f for k, f in spec["API0001"]["request"].items() if f.get("name", k) == "id"]
    assert sorted((f["location"], f["type"]) for f in ids) == [("body", "integer"), ("path", "string")]


def test_matching_rows_have_no_difference(tmp_path: Path):
    path = _excel(tmp_path, [("id", "文字列", "○", 10, "path"), ("id", "数値", "○", None, "body"), ("name", "文字列", "", 20, "body")])
    result = app_spec_diff.run(CONF, path, _oas(tmp_path))
    assert result.changes == []


def test_difference_of_body_field_is_detected_even_if_path_parameter_has_same_name(tmp_path: Path):
    # body の id の型を誤記（OpenAPI は integer）. 以前は path の id と 1 つにまとまり、型の違いを見落とした
    path = _excel(tmp_path, [("id", "文字列", "○", 10, "path"), ("id", "文字列", "○", None, "body"), ("name", "文字列", "", 20, "body")])
    result = app_spec_diff.run(CONF, path, _oas(tmp_path))
    mism = [c for c in result.changes if c.category == "type_mismatch"]
    assert len(mism) == 1 and mism[0].before == "string" and mism[0].after == "integer"
    assert "body" in mism[0].field


def test_missing_body_field_with_same_name_as_path_parameter_is_reported(tmp_path: Path):
    path = _excel(tmp_path, [("id", "文字列", "○", 10, "path"), ("name", "文字列", "", 20, "body")])
    result = app_spec_diff.run(CONF, path, _oas(tmp_path))
    missing = [c for c in result.changes if c.category == "field_missing_in_doc"]
    assert len(missing) == 1 and "id" in missing[0].field and "body" in missing[0].field
    assert not [c for c in result.changes if c.category == "location_mismatch"]


def test_unique_name_with_different_location_is_still_location_mismatch():
    doc = {"A": {"request": {"id": {"type": "string", "location": "query"}}}}
    oas = {"A": {"request": {"id": {"type": "string", "location": "path"}}}}
    changes = spec_diff.compare(doc, oas)
    assert [(c.category, c.field, c.before, c.after) for c in changes] == [("location_mismatch", "id", "query", "path")]
