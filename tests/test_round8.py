"""第 8 回: 設定の形の誤りを終了コード 2 に・改訂で消した表を「設定の見直し」にしない・パス共通パラメーターの上書き・ドキュメントの一致."""

from __future__ import annotations

import contextlib
import io
import re
from pathlib import Path

import pytest
import tomllib
from devtools_common.cli import run_main
from openpyxl import Workbook

from design_diff.application import excel_diff
from design_diff.cli import main as cli
from design_diff.infrastructure import oas_reader

ROOT = Path(__file__).resolve().parent.parent


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue() + err.getvalue()


def _book(path: Path, with_table: bool) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "API0001"
    ws["A1"], ws["B1"] = "API 名", "会員照会"
    if with_table:
        ws["A3"] = "入力項目"
        ws.append(["項目名", "型"])
        ws.append(["id", "string"])
    wb.save(path)
    return path


CONF = """
sheets:
  - name: "API*"
    items:
      - {kind: label, name: API 名, label: API 名}
      - {kind: table, name: 入力項目, title: 入力項目, header: 項目名, key: [項目名]}
"""


def _conf_dir(tmp_path: Path, text: str = CONF, name: str = "t.yaml") -> Path:
    d = tmp_path / "config"
    d.mkdir(exist_ok=True)
    (d / name).write_text(text, encoding="utf-8")
    return d


# --- 1. 設定の形の誤り（offset・sheets の要素の型）は終了コード 2 -------------------------------------


@pytest.mark.parametrize(
    "conf,needle",
    [
        ("sheets: [API]", "sheets[0] はキーと値"),
        ("sheets:\n  - {name: 1, items: []}", "sheets[0].name は文字列"),
        ("sheets:\n  - {items: {a: 1}}", "items は項目の配列"),
        ("sheets:\n  - {items: [label]}", "items[0] はキーと値"),
        ("sheets:\n  - {items: [{kind: label, name: a, label: a, offset: 1}]}", "offset は [行, 列] の整数 2 つ"),
        ("sheets:\n  - {items: [{kind: label, name: a, label: a, offset: [0, x]}]}", "offset は [行, 列] の整数 2 つ"),
        ("sheets:\n  - {items: [{kind: label, name: a, label: a, offset: [0, 1, 2]}]}", "offset は [行, 列] の整数 2 つ"),
        ("sheets:\n  - {items: [{kind: label, name: a, label: 3}]}", "label は文字列"),
        ("sheets:\n  - {items: [{kind: table, name: a, header: [x], key: [x]}]}", "header は文字列"),
        ("- a\n- b", "キーと値で書きます"),
    ],
)
def test_excel_config_shape_errors_are_exit_2(tmp_path: Path, conf: str, needle: str):
    d = _conf_dir(tmp_path, conf)
    b = _book(tmp_path / "b.xlsx", True)
    rc, out = run("--config-dir", str(d), "excel", "--type", "t", "--before", str(b), "--after", str(b))
    assert rc == 2, out
    assert needle in out


def test_valid_offset_is_accepted():
    excel_diff.validate_config({"sheets": [{"items": [{"kind": "label", "name": "a", "label": "a", "offset": [1, 0]}]}]}, "x")


@pytest.mark.parametrize("pair", ["{type: [api], before: a.xlsx, after: b.xlsx}", "{type: 1, before: a.xlsx, after: b.xlsx}"])
def test_all_with_non_string_type_is_exit_2(tmp_path: Path, pair: str):
    d = _conf_dir(tmp_path, f"pairs:\n  - {pair}\n", "global.yaml")
    rc, out = run("--config-dir", str(d), "excel", "--all")
    assert rc == 2, out
    assert "pairs[0] の type は文字列" in out


def test_all_with_pairs_not_a_list_is_exit_2(tmp_path: Path):
    d = _conf_dir(tmp_path, "pairs: {type: api}\n", "global.yaml")
    rc, out = run("--config-dir", str(d), "excel", "--all")
    assert rc == 2 and "pairs は配列" in out


# --- 2. 改訂で表を消した・足したのは差分（設定の見直しではない） ---------------------------------------


def test_removed_table_is_a_diff_not_a_config_problem(tmp_path: Path):
    d = _conf_dir(tmp_path)
    b, a = _book(tmp_path / "b.xlsx", True), _book(tmp_path / "a.xlsx", False)
    rc, out = run("--config-dir", str(d), "excel", "--type", "t", "--before", str(b), "--after", str(a))
    assert rc == 0, out  # 前は table_not_found（設定の見直しが必要）で終了コード 1
    assert "削除された表" in out and "設定の見直し" not in out
    rc, out = run("--config-dir", str(d), "excel", "--type", "t", "--before", str(a), "--after", str(b), "--fail-on-diff", "-f", "json")
    assert rc == 1 and '"table_added"' in out


def test_table_missing_on_both_sides_is_still_a_config_problem(tmp_path: Path):
    d = _conf_dir(tmp_path)
    b = _book(tmp_path / "b.xlsx", False)
    rc, out = run("--config-dir", str(d), "excel", "--type", "t", "--before", str(b), "--after", str(b))
    assert rc == 1 and "どちらにも無い表（設定の見直しが必要）" in out


# --- 3. パス共通のパラメーターと操作のパラメーター --------------------------------------------------


def _read(tmp_path: Path, text: str) -> dict:
    p = tmp_path / "openapi.yaml"
    p.write_text(text, encoding="utf-8")
    return oas_reader.read(p)


def test_operation_parameter_overrides_path_level_parameter(tmp_path: Path):
    spec = _read(
        tmp_path,
        """
openapi: 3.0.3
components:
  parameters:
    Trace: {name: X-Trace-Id, in: header, schema: {type: string, maxLength: 8}}
  pathItems: {}
paths:
  /members/{id}:
    parameters:
      - {name: id, in: path, required: true, schema: {type: string, maxLength: 5}}
      - {name: q, in: query, schema: {type: string}}
      - $ref: '#/components/parameters/Trace'
    get:
      operationId: API0001
      parameters:
        - {name: id, in: path, required: true, schema: {type: integer}}
        - {name: x-trace-id, in: header, required: true, schema: {type: string, maxLength: 16}}
        - {name: id, in: query, schema: {type: string}}
      responses: {'204': {description: ok}}
    post:
      operationId: API0002
      responses: {'204': {description: ok}}
""",
    )
    req = spec["API0001"]["request"]
    assert req["id"]["type"] == "integer" and req["id"]["max_length"] is None  # 操作のもので上書き
    trace = [f for k, f in req.items() if (f.get("name") or k).lower() == "x-trace-id"]
    assert len(trace) == 1 and trace[0]["required"] is True and trace[0]["max_length"] == 16  # header は大文字・小文字を区別しない
    assert sum(1 for k, f in req.items() if (f.get("name") or k) == "id") == 2  # path の id と query の id は別
    assert set(spec["API0002"]["request"]) == {"id", "q", "X-Trace-Id"}  # 共通のものだけ


def test_path_item_ref_is_resolved(tmp_path: Path):
    spec = _read(
        tmp_path,
        """
openapi: 3.1.0
components:
  pathItems:
    Member:
      parameters: [{name: id, in: path, required: true, schema: {type: string}}]
      get: {operationId: API0001, responses: {'204': {description: ok}}}
paths:
  /members/{id}: {$ref: '#/components/pathItems/Member'}
""",
    )
    assert "id" in spec["API0001"]["request"]


def test_parameters_not_a_list_is_input_error(tmp_path: Path):
    with pytest.raises(oas_reader.OpenApiError, match="配列"):
        _read(tmp_path, "openapi: 3.0.3\npaths:\n  /a:\n    parameters: {name: id}\n    get: {operationId: A, responses: {}}\n")


# --- 4. ドキュメントと実際の動きのずれ（機械で確かめる） ----------------------------------------------


def test_readme_python_version_matches_pyproject_and_ci():
    meta = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    need = re.fullmatch(r">=(\d+\.\d+)", meta["project"]["requires-python"]).group(1)
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    assert f"Python {need} 以上" in readme
    assert f"(3, {need.split('.')[1]})" in (ROOT / "setup.sh").read_text(encoding="utf-8")
    assert f"Python {need} 以上" in (ROOT / "setup.cmd").read_bytes().decode("cp932")
    ci = (ROOT / ".github/workflows/ci.yml").read_text(encoding="utf-8")
    assert f'"{need}"' in ci  # 最低の版を CI で確かめている
    assert meta["tool"]["ruff"]["target-version"] == "py" + need.replace(".", "")


def test_every_cli_option_is_in_readme_and_has_help():
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    options = readme.split("## オプション", 1)[1].split("### 終了コード", 1)[0]
    parser = cli.build_parser()
    subs = next(a for a in parser._actions if a.__class__.__name__ == "_SubParsersAction").choices
    for name, sp in subs.items():
        for act in sp._actions:
            if act.dest in ("help",):
                continue
            flag = max(act.option_strings, key=len)
            assert act.help, f"{name} {flag} に --help の説明がない"
            short = min(act.option_strings, key=len)
            assert f"`{flag}" in options or f"`{short} " in options or f"{flag} " in options, f"README に {name} {flag} が無い"
    fmt = next(a for a in subs["excel"]._actions if a.dest == "format")
    assert f"`-f {'｜'.join(fmt.choices)}`" in options or all(c in options for c in fmt.choices)


def test_setup_sh_supports_git_bash_layout():
    sh = (ROOT / "setup.sh").read_text(encoding="utf-8")
    assert "Git Bash" in sh and ".venv/Scripts" in sh and ".venv/bin/python" not in sh
