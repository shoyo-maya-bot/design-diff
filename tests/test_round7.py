"""第 7 回: object 配列の項目・結合したラベル・default 応答・spec の終了コード."""

from __future__ import annotations

import contextlib
import io
from pathlib import Path
from unittest import mock

import pytest
from devtools_common.cli import run_main
from openpyxl import Workbook

from design_diff.application import spec_diff as spec_app
from design_diff.cli import main as cli
from design_diff.infrastructure import excel_reader, excel_spec_reader, oas_reader

ROOT = Path(__file__).resolve().parent.parent


def _oas(tmp_path: Path, text: str) -> Path:
    p = tmp_path / "openapi.yaml"
    p.write_text(text, encoding="utf-8")
    return p


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue() + err.getvalue()


# --- object の配列の項目（orders[].id → orders.id）------------------------------------------------

ARRAY_OAS = """
openapi: 3.0.3
paths:
  /members/{id}:
    get:
      operationId: API0001
      responses:
        '200':
          content:
            application/json:
              schema:
                type: object
                properties:
                  name: {type: string}
                  orders:
                    type: array
                    items: {$ref: '#/components/schemas/Order'}
                  tags:
                    type: array
                    items: {type: string, maxLength: 10}
                  matrix:
                    type: array
                    items:
                      type: array
                      items:
                        type: object
                        required: [x]
                        properties: {x: {type: integer}}
components:
  schemas:
    Order:
      type: object
      required: [id]
      properties:
        id: {type: string, maxLength: 12}
        lines:
          type: array
          items:
            type: object
            properties: {sku: {type: string}}
"""


def test_object_array_items_are_expanded_as_parent_dot_child(tmp_path: Path):
    res = oas_reader.read(_oas(tmp_path, ARRAY_OAS))["API0001"]["response"]
    assert res["orders"]["type"] == "array"
    assert res["orders.id"] == {"type": "string", "required": True, "max_length": 12, "location": "body"}
    assert res["orders.lines"]["type"] == "array"
    assert res["orders.lines.sku"]["type"] == "string"  # 配列の中の配列の object も展開する
    assert res["matrix.x"]["type"] == "integer"  # 配列の配列の object も要素の項目にする
    assert res["tags"] == {"type": "array", "required": False, "max_length": 10, "location": "body"}
    assert not any(k.startswith("tags.") for k in res)  # 文字列の配列は要素の項目を作らない


def test_object_array_item_in_doc_is_not_reported_missing():
    doc = {"API0001": {"method": "GET", "path": "/members/{id}", "request": {}, "statuses": [], "response": {}}}
    doc["API0001"]["response"] = {
        "orders": {"type": "array", "required": False, "max_length": None, "location": "body"},
        "orders.id": {"type": "string", "required": True, "max_length": 12, "location": "body"},
    }
    oas = {"API0001": {"method": "GET", "path": "/members/{x}", "request": {}, "statuses": [], "response": {}}}
    oas["API0001"]["response"] = {
        "orders": {"type": "array", "required": False, "max_length": None, "location": "body"},
        "orders.id": {"type": "string", "required": True, "max_length": 12, "location": "body"},
    }
    from design_diff.domain import spec_diff

    assert spec_diff.compare(doc, oas) == []


def test_recursive_array_schema_stops_with_note(tmp_path: Path):
    text = """
openapi: 3.0.3
paths:
  /t:
    get:
      operationId: T
      responses:
        '200': {content: {application/json: {schema: {$ref: '#/components/schemas/Node'}}}}
components:
  schemas:
    Node:
      type: object
      properties:
        name: {type: string}
        children: {type: array, items: {$ref: '#/components/schemas/Node'}}
"""
    spec, notes = oas_reader.read_with_notes(_oas(tmp_path, text))
    res = spec["T"]["response"]
    assert set(res) == {"name", "children"}
    assert any("循環参照" in n for n in notes)


# --- 結合したラベル -----------------------------------------------------------------------------


def test_label_merged_across_columns_reads_cell_right_of_merge():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "API ID", "API0001"
    ws.merge_cells("A1:B1")  # ラベルを 2 列に結合
    assert excel_reader.read_label(ws, "API ID") == "API0001"  # 以前は結合の中（B1）を読んでラベル自身


def test_label_merged_down_reads_cell_below_merge():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A3"] = "API ID", "API0001"
    ws.merge_cells("A1:A2")
    assert excel_reader.read_label(ws, "API ID", (1, 0)) == "API0001"


def test_block_label_merged_down_reads_rows_below_merge():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A3"], ws["A4"] = "処理概要", "1 行目", "2 行目"
    ws.merge_cells("A1:A2")
    assert excel_reader.read_block(ws, "処理概要") == "1 行目\n2 行目"  # 以前は 1 行目にラベル自身が入った


def test_unmerged_label_offset_is_unchanged():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "API ID", "API0001", "x"
    assert excel_reader.read_label(ws, "API ID") == "API0001"
    assert excel_reader.read_label(ws, "API ID", (0, 2)) == "x"


# --- default 応答 -------------------------------------------------------------------------------


def test_default_response_is_not_a_status(tmp_path: Path):
    text = """
openapi: 3.0.3
paths:
  /x:
    get:
      operationId: API0001
      responses:
        '200': {description: ok}
        '404': {description: nf}
        default: {description: error}
"""
    spec, notes = oas_reader.read_with_notes(_oas(tmp_path, text))
    assert spec["API0001"]["statuses"] == ["200", "404"]
    assert any("default" in n and "API0001" in n for n in notes)

    from design_diff.domain import spec_diff

    doc = {"API0001": {"method": "GET", "path": "/x", "request": {}, "response": {}, "statuses": ["200", "404"]}}
    assert spec_diff.compare(doc, spec) == []  # 以前は status_missing_in_doc（default）


# --- spec の終了コード: 内部の KeyError / ValueError は 3、設定の誤りは 2 ------------------------------


def test_internal_key_error_in_spec_is_exit_3(tmp_path: Path):
    with mock.patch.object(spec_app.spec_diff, "compare", side_effect=KeyError("bug")):
        rc, out = run(
            "--config-dir", str(ROOT / "config"), "spec",
            "--excel", str(ROOT / "samples/spec/API定義書.xlsx"), "--oas", str(ROOT / "samples/spec/openapi.yaml"),
        )  # fmt: skip
    assert rc == 3, out


def test_internal_value_error_in_spec_is_exit_3(tmp_path: Path):
    with mock.patch.object(spec_app.spec_diff, "compare", side_effect=ValueError("bug")):
        rc, out = run(
            "--config-dir", str(ROOT / "config"), "spec",
            "--excel", str(ROOT / "samples/spec/API定義書.xlsx"), "--oas", str(ROOT / "samples/spec/openapi.yaml"),
        )  # fmt: skip
    assert rc == 3, out


@pytest.mark.parametrize(
    "conf",
    [
        "sheets: 'API*'\nmethod: {label: メソッド}\npath: {label: パス}\n",  # api_id が無い
        "sheets: 'API*'\napi_id: API ID\nmethod: {label: メソッド}\npath: {label: パス}\n",  # 辞書でない
        "sheets: 'API*'\napi_id: {label: API ID, offset: [1]}\nmethod: {label: メソッド}\npath: {label: パス}\n",
        "sheets: 'API*'\napi_id: {label: API ID}\nmethod: {label: メソッド}\npath: {label: パス}\nrequest: {title: 入力項目}\n",
        "sheets: 'API*'\napi_id: {label: API ID}\nmethod: {label: メソッド}\npath: {label: パス}\nstatuses: {title: 応答}\n",
        "sheets: 'API*'\napi_id: {label: API ID}\nmethod: {label: メソッド}\npath: {label: パス}\ntype_map: [a]\n",
        "- 1\n",  # 辞書でない
    ],
)
def test_spec_config_errors_are_exit_2(tmp_path: Path, conf: str):
    cfg = tmp_path / "spec.yaml"
    cfg.write_text(conf, encoding="utf-8")
    rc, out = run(
        "spec", "--config", str(cfg),
        "--excel", str(ROOT / "samples/spec/API定義書.xlsx"), "--oas", str(ROOT / "samples/spec/openapi.yaml"),
    )  # fmt: skip
    assert rc == 2, out
    assert "Traceback" not in out
    assert "spec.yaml" in out


def test_sample_spec_config_is_valid():
    from devtools_common.config import load_structured

    excel_spec_reader.validate_config(load_structured(ROOT / "config/spec.yaml"), "spec.yaml")


# --- ドキュメントと実際の動きのずれ（機械で確かめる）----------------------------------------------


def test_readme_exit_code_table_matches_cli():
    import re

    from devtools_common.cli import ExitCode

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    table = readme.split("### 終了コード", 1)[1].strip().split("\n\n", 1)[0]
    codes = {int(m) for m in re.findall(r"^\| `(\d)` \|", table, re.M)}
    assert codes == {int(c) for c in ExitCode}
    doc = cli.__doc__ or ""
    for c in ExitCode:  # --help（モジュールの docstring）にも同じ番号がある
        assert re.search(rf"(?:^|[ /]){int(c)} ", doc.split("終了コード:", 1)[1]), c


def test_every_rule_id_is_documented_in_specification():
    import re

    from design_diff.application import render

    spec = (ROOT / "docs/SPECIFICATION.md").read_text(encoding="utf-8")
    src = "\n".join(p.read_text(encoding="utf-8") for p in (ROOT / "src").rglob("*.py"))
    ids = set(render.LABELS) | set(render.ALERT_LABELS) | set(re.findall(r'Finding\(\s*"([a-z_]+)"', src))
    missing = sorted(i for i in ids if f"`{i}`" not in spec)
    assert missing == []
