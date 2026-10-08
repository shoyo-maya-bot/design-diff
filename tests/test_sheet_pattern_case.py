"""シート名のパターン（sheets / name）は大文字・小文字を区別して照合する（OS で一致が変わらない）.

fnmatch は Windows で大文字・小文字を無視するので、Windows の振る舞い（os.path.normcase が小文字にする）を真似て確かめる.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from design_diff.application import excel_diff
from design_diff.infrastructure import excel_spec_reader


@pytest.fixture
def windows_like_normcase(monkeypatch):
    monkeypatch.setattr(os.path, "normcase", lambda s: s.lower().replace("/", "\\"))


def _book(path: Path, titles: list[str]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for t in titles:
        ws = wb.create_sheet(t)
        ws["A1"], ws["B1"] = "API ID", t.upper()
        ws["A2"], ws["B2"] = "メソッド", "GET"
        ws["A3"], ws["B3"] = "パス", "/x"
    wb.save(path)
    return path


def test_spec_sheet_pattern_is_case_sensitive(tmp_path: Path, windows_like_normcase):
    path = _book(tmp_path / "s.xlsx", ["API1", "api2"])
    conf = {"sheets": "API*", "api_id": {"label": "API ID"}, "method": {"label": "メソッド"}, "path": {"label": "パス"}}
    spec, _, _ = excel_spec_reader.read(path, conf)
    assert list(spec) == ["API1"]


def test_excel_diff_sheet_pattern_is_case_sensitive(tmp_path: Path, windows_like_normcase):
    before = _book(tmp_path / "b.xlsx", ["API1", "api2"])
    after = _book(tmp_path / "a.xlsx", ["API1"])
    conf = {"sheets": [{"name": "API*", "items": []}]}
    result = excel_diff.run(conf, before, after, outside="off")
    assert not [c for c in result.changes if c.category == "sheet_removed"]
