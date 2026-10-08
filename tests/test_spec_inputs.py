"""spec 比較の入力: 壊れた OpenAPI は入力の誤り（終了コード 2）、同じ API ID のシートは最初のシートを使って警告."""

from __future__ import annotations

import contextlib
import io
from pathlib import Path

import pytest
from devtools_common.cli import run_main
from openpyxl import Workbook

from design_diff.application import spec_diff
from design_diff.cli import main as cli

ROOT = Path(__file__).resolve().parent.parent
SPEC_CONF = {
    "sheets": "API*",
    "api_id": {"label": "API ID"},
    "method": {"label": "メソッド"},
    "path": {"label": "パス"},
}


def run(*args: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue() + err.getvalue()


# --- 壊れた OpenAPI は終了コード 2 --------------------------------------------------------------


@pytest.mark.parametrize(
    "content",
    [
        "openapi: 3.0.0\npaths: {\n  /x: [\n",  # YAML の構文の誤り
        "openapi: 3.0.0\npaths:\n  /x: [1, 2]\n",  # パスの中身が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get: 1\n",  # 操作が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get:\n      parameters: [1]\n",  # パラメーターが辞書でない
        "openapi: 3.0.0\npaths: [1]\n",  # paths が辞書でない
        "openapi: 3.0.0\npaths:\n  /x:\n    get:\n      responses: {'200': {content: {application/json: {schema: {$ref: '#/nope'}}}}}\n",
    ],
)
def test_broken_openapi_is_input_error(tmp_path: Path, content: str):
    oas = tmp_path / "openapi.yaml"
    oas.write_text(content, encoding="utf-8")
    rc, out = run("spec", "--excel", str(ROOT / "samples/spec/API定義書.xlsx"), "--oas", str(oas))
    assert rc == 2, out
    assert "Traceback" not in out


def test_broken_openapi_raises_value_error(tmp_path: Path):
    from design_diff.infrastructure import oas_reader

    oas = tmp_path / "o.yaml"
    oas.write_text("paths: {\n", encoding="utf-8")
    with pytest.raises(ValueError, match="OpenAPI"):
        oas_reader.read_with_notes(oas)


# --- 同じ API ID のシート -----------------------------------------------------------------------


def test_duplicate_api_id_uses_first_sheet_and_warns(tmp_path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    for title, method in (("API1", "GET"), ("API2", "POST")):
        ws = wb.create_sheet(title)
        ws["A1"], ws["B1"] = "API ID", "API0001"
        ws["A2"], ws["B2"] = "メソッド", method
        ws["A3"], ws["B3"] = "パス", "/x"
    path = tmp_path / "spec.xlsx"
    wb.save(path)
    oas = tmp_path / "o.yaml"
    oas.write_text("openapi: 3.0.0\npaths:\n  /x:\n    get:\n      operationId: API0001\n      responses: {'200': {}}\n", encoding="utf-8")
    result = spec_diff.run(SPEC_CONF, path, oas)
    alerts = [w for rule, w in result.alerts if rule == "duplicate_api_id"]
    assert len(alerts) == 1 and "「API1」" in alerts[0] and "「API2」" in alerts[0]
    assert not [c for c in result.changes if c.category == "method_mismatch"]  # 最初のシート（GET）を使う
